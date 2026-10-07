# %% [markdown]
# # Caso 2 · Terremoto en Colombia (10-ago-2026)
# ## Notebook 01 — Construcción y validación de la instancia: demanda, candidatos, coordenadas y distancias por carretera
#
# **Universidad de La Sabana · Diseño y Gestión de la Cadena de Suministro · Prof. Gonzalo Mejía**
#
# Este notebook toma los **datos originales del grupo** (`data/raw/`) y produce los **datos procesados** que usa el modelo
# (`data/processed/`). No cambia ninguna decisión del grupo: recalcula cada valor derivado, lo compara con lo que ya estaba en
# los CSV y en el Excel, y se detiene si algo no coincide.
#
# | Sección | Qué hace |
# |---|---|
# | 1. Configuración | Rutas relativas, parámetros de ejecución |
# | 2. Parámetros | Lee la hoja `Parametros` del Excel del grupo (fuente única de supuestos) |
# | 3. Demanda | Verifica población DANE, recalcula la demanda $d_i$ para cada fracción $f$ y la contrasta con la tasa de albergue observada |
# | 4. Candidatos y presupuesto | Recalcula categoría, capacidad, costo fijo y presupuesto |
# | 5. Coordenadas | Cabeceras municipales DIVIPOLA (DANE, datos.gov.co) |
# | 6. Distancias por carretera | Matriz OSRM (OpenStreetMap) 29 × 18 + verificación Google Maps de 164 pares → matriz final |
# | 7. Validación de la matriz | Razón carretera/línea recta, alcance a 180 km y criterio de 'ciudad cercana', OSRM vs. Google Maps, pares cerca del corte, verificación 140–160 km |
# | 8. Salidas | CSV procesados, base de datos Excel (LEEME, Fuentes, Diccionario…), `docs/bibliografia.md`, `docs/diccionario_datos.md` y figura de nodos |
#
# **Cómo ejecutarlo:** desde la carpeta `notebooks/` con *Run All*. No necesita internet: usa las respuestas de DIVIPOLA y OSRM
# guardadas en `data/raw/`. Para volver a consultar los servicios, cambie `ACTUALIZAR_COORDENADAS` / `ACTUALIZAR_OSRM` a `True`.

# %% [markdown]
# ## 1. Configuración

# %%
import json, math, re, shutil
from pathlib import Path
import numpy as np
import pandas as pd
import matplotlib
import matplotlib.pyplot as plt
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from IPython.display import display, Markdown

ROOT = Path.cwd().parent if Path.cwd().name == "notebooks" else Path.cwd()
RAW, PROC, FIG = ROOT / "data" / "raw", ROOT / "data" / "processed", ROOT / "results" / "figuras"
TAB = ROOT / "results" / "tablas"
import sys
sys.path.insert(0, str(ROOT / "tools"))   # módulo de calidad de datos (tools/calidad_datos.py)
from calidad_datos import evaluar as evaluar_calidad, resumen as resumen_calidad
for p in (PROC, FIG, TAB):
    p.mkdir(parents=True, exist_ok=True)

ACTUALIZAR_COORDENADAS = False   # True: vuelve a descargar DIVIPOLA de datos.gov.co
ACTUALIZAR_OSRM = False          # True: vuelve a consultar la matriz en router.project-osrm.org
OSRM = "https://router.project-osrm.org"
HEADERS = {"User-Agent": "CursoDiseno-CadenaSuministro-LaSabana/1.0 (trabajo academico)"}
BBOX = dict(lat=(1.5, 7.0), lon=(-78.5, -74.0))   # Valle del Cauca + Eje Cafetero

def md(t): display(Markdown(t))
def es(x, d=0):
    """Formato numérico colombiano: punto de miles y coma decimal."""
    return f"{x:,.{d}f}".replace(",", "X").replace(".", ",").replace("X", ".")
def redondear(v):
    """Redondeo 'mitad hacia arriba', igual que ROUND de Excel (numpy redondea al par)."""
    return int(math.floor(v + 0.5))

print("Raíz del proyecto:", ROOT.name, "| carpetas:", ", ".join(sorted(p.name for p in ROOT.iterdir() if p.is_dir() and not p.name.startswith("."))))

# %% [markdown]
# ## 2. Parámetros del grupo
# La hoja `Parametros` de `base_datos_caso2_original.xlsx` es la **fuente única** de supuestos. Se leen los valores de entrada
# (columna C) con su tipo (Real / Supuesto) y fuente.

# %%
wb0 = load_workbook(RAW / "base_datos_caso2_original.xlsx")
ws = wb0["Parametros"]
par_tab = pd.DataFrame([[ws.cell(r, c).value for c in range(1, 7)] for r in range(5, ws.max_row + 1)],
                       columns=["ID", "Parametro", "Valor", "Unidad", "Tipo", "Fuente"]).dropna(subset=["ID"])
P = dict(zip(par_tab.ID, par_tab.Valor))
display(par_tab.style.hide(axis="index"))

C_F, T_MESES, V_KIT, TARIFA = P["P01"], P["P02"], P["P03"], P["P04"]
FACTOR = {"Grande": P["P05"], "Intermedia": P["P06"], "Pequeña": P["P07"]}
CAP = {"Grande": P["P08"], "Intermedia": P["P09"], "Pequeña": P["P10"]}
CORTE_G, CORTE_I = P["P11"], P["P12"]
D_MAX, FRAC_B = P["P16"], P["P17"]
RED_A, RED_B = P["P18"], P["P19"]
F_FRAC = {"d_f02": P["P21"], "d_f05": P["P22"], "d_f10": P["P20"], "d_f13": P["P23"]}
HOGAR_ALT = P["P24"]

# Costo variable v (P03) separado en sus dos componentes explícitos (enunciado: 1 kit de alimentación + 1 kit de aseo por persona).
# Precios de la fuente del grupo (F06, El Tiempo 14-ago-2026, precios minoristas D1): 29.730 + 29.730 = 59.460 COP.
# El modelo usa v = 60.000 COP (P03): es un REDONDEO declarado de 59.460 (+540 COP/persona, +0,9 %), no un precio distinto.
KIT_ALIMENTACION = 29_730   # COP/persona — kit de alimentación
KIT_ASEO = 29_730           # COP/persona — kit de aseo
AJUSTE_REDONDEO_KITS = V_KIT - (KIT_ALIMENTACION + KIT_ASEO)
assert V_KIT == 60_000 and AJUSTE_REDONDEO_KITS == 540, "P03 debe ser el redondeo de 29.730 + 29.730"
md(f"**Costo variable por persona:** kit de alimentación {es(KIT_ALIMENTACION)} + kit de aseo {es(KIT_ASEO)} = {es(KIT_ALIMENTACION + KIT_ASEO)} COP; "
   f"el modelo usa **v = {es(V_KIT)} COP** (P03), redondeo declarado de +{es(AJUSTE_REDONDEO_KITS)} COP/persona "
   f"(+{es(100 * AJUSTE_REDONDEO_KITS / (KIT_ALIMENTACION + KIT_ASEO), 1)} %).")

# %% [markdown]
# ## 3. Demanda: 29 municipios afectados
# **Construcción (decisión del grupo):** $d_i = \text{round}(NH_i \cdot h_i \cdot f)$, donde $NH_i$ son las viviendas no habitables
# del RUD, $h_i$ = personas registradas / familias registradas (tamaño de hogar RUD del municipio) y $f$ la fracción que requiere
# campamento (base 10 %). Se incluyen los municipios con $NH_i \ge 400$ (filtro S1).
#
# Primero se verifica la población contra el archivo oficial del DANE (PPED 2018-2042, actualizado el 30-jul-2025).

# %%
dem = pd.read_csv(RAW / "demanda_terremoto.csv", dtype={"divipola": str})
cand = pd.read_csv(RAW / "candidatos_costos.csv", dtype={"divipola": str})

pped = pd.read_excel(RAW / "dane" / "PPED-AreaMun-2018-2042_VP.xlsx", sheet_name="PobMunicipalxÁrea", header=7, dtype={"MPIO": str})
pped = pped[(pped["AÑO"] == 2026) & (pped["ÁREA GEOGRÁFICA"] == "Total")].set_index("MPIO")["TOTAL"]
dcd = pd.read_excel(RAW / "dane" / "DCD-area-proypoblacion-Mun-2020-2035-ActPostCOVID-19.xlsx", sheet_name="Hoja1", header=8, dtype={"MPIO": str})
dcd = dcd[(dcd["AÑO"] == 2026) & (dcd["ÁREA GEOGRÁFICA"] == "Total")].set_index("MPIO")["Población"]

control_pob = pd.concat([dem[["divipola", "municipio", "pob_DANE_2026"]].assign(rol="demanda"),
                         cand[["divipola", "municipio", "pob_DANE_2026"]].assign(rol="candidato")], ignore_index=True)
control_pob["PPED_2026"] = control_pob.divipola.map(pped)
control_pob["DCD_postCOVID_2026"] = control_pob.divipola.map(dcd)
control_pob["coincide_PPED"] = control_pob.pob_DANE_2026 == control_pob.PPED_2026
assert control_pob.coincide_PPED.all(), control_pob[~control_pob.coincide_PPED]
md(f"✔ **Población:** los {len(control_pob)} municipios coinciden exactamente con DANE PPED 2026.")

# Recalcular la demanda y comparar con el CSV del grupo
h = dem.personas_RUD / dem.familias_RUD
calc = pd.DataFrame({"hogar_RUD": h.round(2), "personas_en_NH": (dem.NH_RUD * h).apply(redondear)})
for col, f in F_FRAC.items():
    calc[col] = (dem.NH_RUD * h * f).apply(redondear)
calc["d_f10_hogar3"] = (dem.NH_RUD * HOGAR_ALT * F_FRAC["d_f10"]).apply(redondear)
calc["pct_registrado"] = (dem.personas_RUD / dem.pob_DANE_2026 * 100).round(1)
dif = (calc - dem[calc.columns]).abs().max()
display(dif.to_frame("diferencia máxima vs CSV del grupo"))
assert (dif == 0).all(), "La demanda recalculada no coincide con el CSV del grupo"
assert (dem.NH_RUD >= P["P15"]).all() and {"Cali", "Pereira"} <= set(dem.municipio) and len(dem) >= 20
md(f"✔ **Demanda:** {len(dem)} municipios (incluye Cali y Pereira), todos con NH ≥ {int(P['P15'])}. "
   f"Demanda total: f = 2 % → {es(dem.d_f02.sum())} · 5 % → {es(dem.d_f05.sum())} · **10 % (base) → {es(dem.d_f10.sum())}** · "
   f"13 % → {es(dem.d_f13.sum())} · hogar 3,0 → {es(dem.d_f10_hogar3.sum())} personas.")

# %%
resumen_dep = dem.groupby("departamento").agg(municipios=("municipio", "count"), NH=("NH_RUD", "sum"),
                                              personas_en_NH=("personas_en_NH", "sum"), demanda_base=("d_f10", "sum"))
resumen_dep.loc["TOTAL"] = resumen_dep.sum()
display(resumen_dep.style.format("{:,.0f}"))
md("**Calidad del dato (señalado para validación):** la razón familias/NH varía entre "
   f"{es((dem.familias_RUD/dem.NH_RUD).min(),1)} y {es((dem.familias_RUD/dem.NH_RUD).max(),1)} "
   f"(máximo en {dem.loc[(dem.familias_RUD/dem.NH_RUD).idxmax(),'municipio']}). Como la demanda usa $NH \\times h$ "
   "(una familia por vivienda no habitable), municipios con muchas familias registradas por vivienda podrían estar subestimados. "
   "Los datos RUD provienen de un agregador secundario (datosdelterremoto.org, corte 17-sep-2026).")

# %% [markdown]
# ### 3.1 Tasa de albergue observada vs. fracción de planeación $f$
# El grupo usa $f$ = 10 % como **escenario de planeación**: es cercano al 11,6 % que estima ABAG (2017, metodología Hazus) para el
# escenario Hayward, con rango 8,3–13,2 % por condado (F10). Ese porcentaje es la fracción de **desplazados** que busca refugio
# público y está calibrado con demografía de EE. UU. Para contrastarlo con lo observado en este sismo se usan las cifras publicadas de
# personas en albergues (`data/raw/albergados_observados.csv`, fuentes F14 y F15) y se divide entre las personas en viviendas no
# habitables de cada municipio (`personas_en_NH` de la demanda).

# %%
alb = pd.read_csv(RAW / "albergados_observados.csv")
alb = alb.merge(dem[["municipio", "personas_en_NH", "d_f10"]], on="municipio", how="left", validate="one_to_one")
assert alb.personas_en_NH.notna().all(), "Hay un municipio con albergados que no está en la demanda"
alb["tasa_observada"] = alb.albergados / alb.personas_en_NH
alb["veces_f_base"] = F_FRAC["d_f10"] / alb.tasa_observada
display(alb[["municipio", "albergados", "fecha_dato", "tipo_dato", "personas_en_NH", "tasa_observada", "veces_f_base", "fuente"]]
        .sort_values("tasa_observada").style.format({"albergados": es, "personas_en_NH": es, "tasa_observada": lambda v: es(100 * v, 1) + " %",
                                                     "veces_f_base": lambda v: es(v, 1)}).hide(axis="index"))
TASA_MIN, TASA_MAX = alb.tasa_observada.min(), alb.tasa_observada.max()
md(f"**Tasa observada = albergados ÷ personas en viviendas no habitables:** entre **{es(100 * TASA_MIN, 1)} %** "
   f"({alb.loc[alb.tasa_observada.idxmin(), 'municipio']}) y **{es(100 * TASA_MAX, 1)} %** ({alb.loc[alb.tasa_observada.idxmax(), 'municipio']}). "
   "**Las fechas no son comparables:** cuatro cifras son un corte puntual del 18-ago-2026 (día 8) y la de Dosquebradas es el promedio "
   "de los primeros 48 días de la contingencia (el 27-sep quedaban 45 personas). Además, quien no está en un albergue oficial puede "
   "estar en casa de familiares, en arriendo o en alojamientos no reportados, así que la tasa observada no mide toda la necesidad. "
   f"Por eso **f = 10 % es un escenario de planeación, no una tasa observada**: equivale a entre {es(alb.veces_f_base.min(), 1)} y "
   f"{es(alb.veces_f_base.max(), 1)} veces lo observado. La sensibilidad del notebook 02 cubre f = 2 % y 5 % (rango observado) y "
   "8,3 %, 11,6 % y 13 % (rango ABAG).")

# %% [markdown]
# ## 4. Candidatos, costos y presupuesto
# **Reglas del grupo:** categoría por población (Grande ≥ 200.000; Intermedia ≥ 100.000; Pequeña en otro caso), capacidad
# 3.000 / 1.000 / 500, costo fijo $F_j = K_j \cdot c_f \cdot T \cdot \phi_{cat}$ y costo variable de un kit de alimentación y un kit
# de aseo por persona. Filtros S2/S7: NH < 250 y % de población registrada < 5 %. La Tebaida es **reserva** (no cumple S7).

# %%
def categoria(pob):
    return "Grande" if pob >= CORTE_G else ("Intermedia" if pob >= CORTE_I else "Pequeña")

c2 = cand.copy()
c2["categoria_calc"] = c2.pob_DANE_2026.apply(categoria)
c2["K_calc"] = c2.categoria_calc.map(CAP)
c2["fj_calc"] = c2.K_calc * C_F * T_MESES * c2.categoria_calc.map(FACTOR)
c2["filtro_calc"] = np.where((c2.NH_RUD < P["P13"]) & (c2.pct_registrado / 100 < P["P14"]), "Cumple", "No cumple")
for a, b in (("categoria", "categoria_calc"), ("capacidad_K", "K_calc"), ("costo_fijo_fj", "fj_calc"), ("filtros_S2_S7", "filtro_calc")):
    assert (c2[a] == c2[b]).all(), f"{a} no coincide"
assert (c2.costo_variable_vj == V_KIT).all()
assert not set(cand.divipola) & set(dem.divipola), "Un candidato está en un municipio afectado"
md("✔ **Candidatos:** categoría, capacidad, costo fijo, costo variable y filtros coinciden con el CSV del grupo; "
   "ningún candidato está en un municipio de demanda.")

base = cand[cand.rol == "Candidato"]
top10 = base.sort_values(["capacidad_K", "id"], ascending=[False, True]).head(10)
B_BASE = FRAC_B * top10.costo_fijo_fj.sum()
display(top10[["id", "municipio", "categoria", "capacidad_K", "costo_fijo_fj"]].style.format({"capacidad_K": "{:,}", "costo_fijo_fj": "${:,.0f}"}).hide(axis="index"))
md(f"**Presupuesto base** = 50 % × {es(top10.costo_fijo_fj.sum()/1e6,1)} M = **{es(B_BASE/1e6,1)} M COP**. "
   f"Reducción −{RED_A:.0%}: {es(B_BASE*(1-RED_A)/1e6,1)} M · reducción −{RED_B:.0%}: {es(B_BASE*(1-RED_B)/1e6,1)} M. "
   f"Abrir los 17 candidatos costaría {es(base.costo_fijo_fj.sum()/1e6)} M solo en costos fijos "
   f"({es(base.costo_fijo_fj.sum()/B_BASE,1)} veces el presupuesto) para {es(base.capacidad_K.sum())} plazas.")
md("**Nota:** las 10 pequeñas tienen el mismo costo fijo, de modo que la elección de cuáles 3 entran al 'top 10' no altera $B$.")

# Sensibilidad de la categoría a la fuente de población (se documenta; no cambia la decisión del grupo)
c2["categoria_si_DCD"] = c2.divipola.map(dcd).apply(categoria)
cambia = c2[c2.categoria != c2.categoria_si_DCD][["municipio", "pob_DANE_2026", "categoria", "categoria_si_DCD"]]
c2["pob_DCD"] = c2.divipola.map(dcd)
md("**Sensibilidad a la fuente de población:** con la proyección DCD post-COVID (2020-2035) cambiaría la categoría de: "
   + (", ".join(f"{r.municipio} ({r.categoria} → {r.categoria_si_DCD})" for r in cambia.itertuples()) or "ninguno")
   + ". Se mantiene PPED (versión más reciente del DANE). Jamundí (196.875 hab.) está a 1,6 % del corte de 200.000.")

# %% [markdown]
# ## 5. Coordenadas: cabeceras municipales DIVIPOLA
# Fuente: DANE — DIVIPOLA, *Códigos de municipios geolocalizados* (datos.gov.co, recurso `gdxc-w37w`). Las coordenadas se
# consultaron el 30-sep-2026 para los 4 departamentos y se guardaron en `data/raw/divipola_coordenadas_datosgov.csv`
# (formato original con coma decimal).
#
# **Limitación (cabeceras municipales).** Cada municipio se representa con **un solo punto**: el de la cabecera municipal que
# publica DIVIPOLA. Ese punto no es necesariamente la plaza principal (en Cali queda unos 4 km al sur del centro y en Pereira unos
# 3 km al suroccidente) y no representa la población rural ni a los damnificados de corregimientos alejados (Buenaventura, Dagua,
# El Cairo, Argelia tienen gran parte de su territorio lejos de la cabecera). El error es pequeño frente al límite de 180 km, pero
# en trayectos cortos (< 40 km) puede ser del orden de 5–10 %. Se mantiene porque el enunciado pide un punto por municipio y es la
# referencia oficial y reproducible.

# %%
if ACTUALIZAR_COORDENADAS:
    import requests
    filas = []
    for dp in ("17", "63", "66", "76"):
        r = requests.get(f"https://www.datos.gov.co/resource/gdxc-w37w.json?cod_dpto={dp}", headers=HEADERS, timeout=60)
        r.raise_for_status()
        filas += [{k: x[k] for k in ("cod_mpio", "nom_mpio", "latitud", "longitud")} for x in r.json()]
    pd.DataFrame(filas).to_csv(RAW / "divipola_coordenadas_datosgov.csv", index=False)

div = pd.read_csv(RAW / "divipola_coordenadas_datosgov.csv", dtype=str)
div["lat"] = div.latitud.str.replace(",", ".").astype(float)
div["lon"] = div.longitud.str.replace(",", ".").astype(float)
div = div.set_index("cod_mpio")

nodos = pd.concat([
    dem[["divipola", "municipio", "departamento"]].assign(tipo="demanda", id=lambda d: "D" + d.divipola, rol="Afectado"),
    cand[["divipola", "municipio", "departamento", "id", "rol"]].assign(tipo="candidato")], ignore_index=True)
nodos["lat"] = nodos.divipola.map(div.lat)
nodos["lon"] = nodos.divipola.map(div.lon)
nodos["fuente_coord"] = "DIVIPOLA (DANE) datos.gov.co gdxc-w37w, consulta 30-sep-2026"
assert nodos[["lat", "lon"]].notna().all().all(), nodos[nodos.lat.isna()]
assert nodos.lat.between(*BBOX["lat"]).all() and nodos.lon.between(*BBOX["lon"]).all()
assert not nodos.duplicated(["lat", "lon"]).any() and nodos.id.is_unique
md(f"✔ **Coordenadas:** {len(nodos)} nodos ({(nodos.tipo=='demanda').sum()} de demanda y {(nodos.tipo=='candidato').sum()} sitios), "
   "sin faltantes, dentro de la región y sin duplicados.")
display(nodos.head(6))

# %% [markdown]
# ## 6. Distancias por carretera (OSRM / OpenStreetMap + verificación en Google Maps)
# **Método.** Servicio `table` de OSRM (perfil *driving*) sobre la red vial de OpenStreetMap. Origen = cabecera de cada municipio
# afectado; destino = cabecera de cada sitio. OSRM ajusta cada punto a la vía más cercana y devuelve la distancia en metros de la
# **ruta más rápida** (no necesariamente la más corta). La consulta se hizo el 30-sep-2026 en 58 llamadas (cada origen contra 9 + 9
# sitios, con coordenadas codificadas en *polyline* de 5 decimales); las respuestas originales y sus URL están en
# `data/raw/osrm/osrm_table_respuestas.csv`. Control de transcripción: se re-consultaron 5 llamadas al azar (50 valores) y
# coincidieron exactamente.
#
# Con `ACTUALIZAR_OSRM = True` la matriz se vuelve a consultar en **una sola llamada** (47 coordenadas) y se compara con la guardada.
#
# **Verificación con Google Maps (1-oct-2026).** Se consultaron en Google Maps (modo carro, primera ruta sugerida) **164 pares**:
# los **139 pares cuya distancia OSRM está entre 160 y 230 km** (la franja donde se decide si un par cumple el límite de 180 km),
# **16 pares de control** (rutas cortas y largas usadas por el modelo y todos los accesos a La Dorada desde Risaralda y Caldas) y
# **9 rutas de montaña de más de 230 km** hacia La Dorada, Aguadas, Supía y Pensilvania. Los resultados y la URL de cada consulta
# están en `data/raw/google_maps/verificacion_google_maps.csv`.
#
# **Regla de la matriz final:** para los 164 pares verificados se usa el valor de Google Maps (fuente que el enunciado nombra primero y
# que el profesor puede repetir); para los 358 pares restantes se usa OSRM. La concordancia en los pares verificados (mediana
# Google/OSRM ≈ 1,005) respalda usar OSRM donde no hubo verificación manual.

# %%
D_nodes = nodos[nodos.tipo == "demanda"].reset_index(drop=True)
S_nodes = nodos[nodos.tipo == "candidato"].reset_index(drop=True)

resp = pd.read_csv(RAW / "osrm" / "osrm_table_respuestas.csv")
M = np.full((len(D_nodes), len(S_nodes)), np.nan)
for r in resp.itertuples():
    v = [float(x) for x in r.distancias_m.split(",")]
    assert v[0] == 0 and len(v) == 10, "Cada respuesta debe traer el origen (0 m) y 9 destinos"
    M[r.origen_idx, 9 * r.bloque: 9 * r.bloque + 9] = np.array(v[1:]) / 1000
dist_km = pd.DataFrame(M, index=D_nodes.id, columns=S_nodes.id)

if ACTUALIZAR_OSRM:
    import requests
    todos = pd.concat([D_nodes, S_nodes], ignore_index=True)
    coords = ";".join(f"{lo:.5f},{la:.5f}" for lo, la in zip(todos.lon, todos.lat))
    params = {"sources": ";".join(map(str, range(len(D_nodes)))),
              "destinations": ";".join(map(str, range(len(D_nodes), len(todos)))), "annotations": "distance,duration"}
    j = requests.get(f"{OSRM}/table/v1/driving/{coords}", params=params, headers=HEADERS, timeout=120).json()
    nueva = pd.DataFrame(j["distances"], index=D_nodes.id, columns=S_nodes.id) / 1000
    dif = (nueva - dist_km).abs()
    print(f"Diferencia máxima con la matriz guardada: {dif.max().max():.2f} km (cambios de OSM desde la consulta original)")
    dist_km = nueva
    pd.DataFrame(j["durations"], index=D_nodes.id, columns=S_nodes.id).div(60).round(1).to_csv(PROC / "tiempos_min.csv")

assert dist_km.notna().all().all(), "Hay pares sin ruta"
dist_osrm = dist_km.copy()

# Verificación Google Maps y matriz final
gm = pd.read_csv(RAW / "google_maps" / "verificacion_google_maps.csv")
assert not gm.duplicated(["id_demanda", "id_sitio"]).any() and gm.km_Google_Maps.notna().all()
for r in gm.itertuples():
    assert abs(dist_osrm.loc[r.id_demanda, r.id_sitio] - r.km_OSRM) < 1e-3, "km_OSRM de la verificación no coincide con la matriz cruda"
fuente_km = pd.DataFrame("OSRM", index=dist_km.index, columns=dist_km.columns)
for r in gm.itertuples():
    dist_km.loc[r.id_demanda, r.id_sitio] = r.km_Google_Maps
    fuente_km.loc[r.id_demanda, r.id_sitio] = "Google Maps"
md(f"Matriz {dist_km.shape[0]} × {dist_km.shape[1]} sin celdas vacías: **{(fuente_km == 'Google Maps').values.sum()} pares con Google Maps** "
   f"y {(fuente_km == 'OSRM').values.sum()} con OSRM.")
display(dist_km.rename(index=dict(zip(D_nodes.id, D_nodes.municipio)), columns=dict(zip(S_nodes.id, S_nodes.municipio)))
        .round(0).head(8))

# %% [markdown]
# ## 7. Validación de la matriz
# 1. **Razón carretera / línea recta** (haversine): debe ser ≥ 1; valores > 2,3 se revisan.
# 2. **Alcance a 180 km:** cada municipio debe tener al menos un candidato a ≤ 180 km y deben existir ≥ 15 candidatos útiles.
# 3. **OSRM vs. Google Maps** en los 164 pares verificados: diferencias y cambios de factibilidad a 180 km.

# %%
def haversine(lat1, lon1, lat2, lon2):
    R = 6371.0088
    p1, p2 = np.radians(lat1), np.radians(lat2)
    a = np.sin((p2 - p1) / 2) ** 2 + np.cos(p1) * np.cos(p2) * np.sin(np.radians(lon2 - lon1) / 2) ** 2
    return 2 * R * np.arcsin(np.sqrt(a))

lin = pd.DataFrame(haversine(D_nodes.lat.values[:, None], D_nodes.lon.values[:, None], S_nodes.lat.values[None, :], S_nodes.lon.values[None, :]),
                   index=D_nodes.id, columns=S_nodes.id)
razon = dist_km / lin
print("Razón carretera / línea recta (matriz final):")
print(razon.stack().describe().round(2).to_string())
assert (razon >= 1.0).all().all(), "Distancia por carretera menor que la línea recta"
nom = dict(zip(nodos.id, nodos.municipio))
altos = razon.stack()[razon.stack() > 2.3].sort_values(ascending=False)
md(f"**{len(altos)} pares con razón > 2,3** (todos de montaña; {sum('C16' in k[1] for k in altos.index)} hacia Pensilvania, "
   "cuyo acceso desde el Eje Cafetero exige rodear la cordillera Central). Ninguno es menor que la línea recta.")
display(pd.DataFrame([(nom[i], nom[j], round(dist_km.loc[i, j], 1), round(lin.loc[i, j], 1), round(v, 2)) for (i, j), v in altos.items()],
                     columns=["origen", "destino", "km carretera", "km línea recta", "razón"]))

# %%
a = (dist_km <= D_MAX).astype(int)
base_ids = list(S_nodes[S_nodes.rol == "Candidato"].id)
por_demanda = pd.DataFrame({"municipio": D_nodes.municipio.values, "departamento": D_nodes.departamento.values,
                            "candidatos_a_180km": a[base_ids].sum(axis=1).values,
                            "dist_min_km": dist_km[base_ids].min(axis=1).round(1).values,
                            "candidato_mas_cercano": dist_km[base_ids].idxmin(axis=1).map(nom).values}, index=D_nodes.id)
por_sitio = pd.DataFrame({"municipio": S_nodes.municipio.values, "rol": S_nodes.rol.values,
                          "demandas_a_180km": a.sum(axis=0).values,
                          "demanda_base_a_180km": (a.mul(dem.set_index("D" + dem.divipola).d_f10, axis=0)).sum(axis=0).values,
                          "dist_min_km": dist_km.min(axis=0).round(1).values}, index=S_nodes.id)
# Criterio explícito de "ciudad cercana" (enunciado: candidatos en ciudades intermedias o grandes "cercanas"):
# un sitio es cercano si al menos un municipio afectado está a <= 180 km por carretera (matriz final). Se reporta también cuántos
# de esos pares solo cumplen gracias a un valor de Google Maps (con OSRM estarían por encima de 180 km).
a_osrm = (dist_osrm <= D_MAX).astype(int)
por_sitio["municipios_a_180km"] = [", ".join(nom[i] for i in dist_km.index if a.loc[i, j]) for j in S_nodes.id]
por_sitio["pares_que_dependen_de_Google"] = [int(((a[j] == 1) & (a_osrm[j] == 0)).sum()) for j in S_nodes.id]
por_sitio["cumple_criterio_ciudad_cercana"] = np.where(por_sitio.demandas_a_180km >= 1, "Sí", "No")
por_sitio["criterio_ciudad_cercana"] = f"≥ 1 municipio afectado a ≤ {D_MAX:.0f} km por carretera (matriz final)"
utiles = por_sitio[(por_sitio.rol == "Candidato") & (por_sitio.demandas_a_180km > 0)]
inutiles = por_sitio[(por_sitio.rol == "Candidato") & (por_sitio.demandas_a_180km == 0)]
assert (por_demanda.candidatos_a_180km > 0).all(), "Hay municipios sin ningún candidato a 180 km"
assert len(utiles) >= 15, "Menos de 15 candidatos útiles"
md(f"✔ Todos los municipios tienen al menos un candidato a ≤ 180 km (mínimo: "
   f"{por_demanda.candidatos_a_180km.min()} en {', '.join(por_demanda[por_demanda.candidatos_a_180km==por_demanda.candidatos_a_180km.min()].municipio)}).  \n"
   f"✔ **{len(utiles)} de 17 candidatos** alcanzan al menos un municipio a ≤ 180 km (mínimo exigido: 15).  \n"
   + (f"⚠️ Sin ningún municipio a ≤ 180 km: **{', '.join(inutiles.municipio)}** (distancia mínima "
      f"{', '.join(f'{d:.0f} km' for d in inutiles.dist_min_km)}).  \n" if len(inutiles) else
      "✔ Ningún candidato queda sin municipios a ≤ 180 km (con la matriz final, La Dorada alcanza a Manizales y Villamaría).  \n") +
   f"Pares factibles (≤ 180 km): {int(a[base_ids].values.sum())} de {a[base_ids].size}.")
display(por_sitio.sort_values("demandas_a_180km").drop(columns="criterio_ciudad_cercana"))
pocos = por_sitio[(por_sitio.rol == "Candidato") & (por_sitio.demandas_a_180km <= 2)]
md(f"**Criterio de 'ciudad cercana' (definición del grupo):** un candidato es cercano si **al menos un municipio afectado está a "
   f"≤ {D_MAX:.0f} km por carretera** (la misma distancia máxima del enunciado). Lo cumplen "
   f"{(por_sitio[por_sitio.rol == 'Candidato'].cumple_criterio_ciudad_cercana == 'Sí').sum()} de 17 candidatos. Los que quedan más justos: "
   + "; ".join(f"**{r.municipio}** alcanza {r.demandas_a_180km} municipio(s) ({r.municipios_a_180km})"
               + (f", y {r.pares_que_dependen_de_Google} de esos pares dependen de un valor de Google Maps" if r.pares_que_dependen_de_Google else "")
               for r in pocos.itertuples())
   + ". La Dorada depende por completo de dos valores de Google Maps (Manizales 166 km y Villamaría 165 km); con OSRM no cumpliría.")

# %%
# OSRM vs. Google Maps en los pares verificados
control = gm.copy()
control["km_linea_recta"] = [round(lin.loc[i, j], 1) for i, j in zip(control.id_demanda, control.id_sitio)]
control["diferencia_km"] = (control.km_Google_Maps - control.km_OSRM).round(1)
control["diferencia_pct"] = (100 * (control.km_Google_Maps / control.km_OSRM - 1)).round(1)
control["factible_OSRM"] = (control.km_OSRM <= D_MAX).astype(int)
control["factible_Google"] = (control.km_Google_Maps <= D_MAX).astype(int)
control["cambia_factibilidad"] = (control.factible_OSRM != control.factible_Google).astype(int)
control["revisar_>10%"] = (control.diferencia_pct.abs() > 10).astype(int)
r_gm = control.km_Google_Maps / control.km_OSRM
md(f"**Google Maps / OSRM** en {len(control)} pares: mediana {r_gm.median():.3f}, percentil 10–90 % "
   f"{r_gm.quantile(.1):.3f}–{r_gm.quantile(.9):.3f}. **{int(control['revisar_>10%'].sum())} pares difieren más de 10 %** y "
   f"**{int(control.cambia_factibilidad.sum())} cambian de factibilidad a 180 km**; en todos ellos manda el valor de Google Maps.")
display(control.loc[(control.cambia_factibilidad == 1) | (control["revisar_>10%"] == 1),
                    ["municipio_origen", "sitio_destino", "km_OSRM", "km_Google_Maps", "diferencia_pct", "factible_OSRM", "factible_Google"]])
c17 = control[control.id_sitio == "C17"].sort_values("km_Google_Maps")
md("**La Dorada (C17).** OSRM daba 212,9 km desde Villamaría y 213,8 km desde Manizales (la ruta 'más rápida' de OSRM rodea por el "
   "sur); Google Maps encuentra la vía Manizales–Letras–Fresno–Mariquita–Honda (Ruta 50) con "
   + ", ".join(f"{r.municipio_origen} {es(r.km_Google_Maps)} km" for r in c17.head(3).itertuples())
   + ". Calculadoras de ruta independientes reportan 170–172 km para Manizales–La Dorada. **Conclusión:** La Dorada sí alcanza "
   "a Manizales y Villamaría a ≤ 180 km; la marca de 'candidato inútil' venía de un error de ruteo de OSRM.")

# %% [markdown]
# **Pares que cambian de factibilidad y cercanía al corte.** Google Maps muestra **kilómetros enteros desde 100 km**, de modo que un
# valor como 178 km puede ser 177,5–178,4 km; además, la ruta "más rápida" cambia con el tráfico. Los pares que quedan a pocos km del
# corte de 180 km son, por eso, frágiles.

# %%
cambian = control[control.cambia_factibilidad == 1].assign(distancia_al_corte_km=lambda d_: (d_.km_Google_Maps - D_MAX).abs())
cambian = cambian.sort_values("distancia_al_corte_km")[["municipio_origen", "sitio_destino", "km_OSRM", "km_Google_Maps",
                                                         "factible_OSRM", "factible_Google", "distancia_al_corte_km"]]
display(cambian.style.format({"km_OSRM": lambda v: es(v, 1), "km_Google_Maps": es, "distancia_al_corte_km": es}).hide(axis="index"))
cerca_corte = cambian[cambian.distancia_al_corte_km <= 3]
md(f"**{len(cerca_corte)} de los {len(cambian)} pares que cambian de factibilidad están a ≤ 3 km del corte:** "
   + ", ".join(f"{r.municipio_origen}–{r.sitio_destino} ({es(r.km_Google_Maps)} km)" for r in cerca_corte.itertuples())
   + ". Su efecto sobre los escenarios oficiales se comprueba en el notebook 02 (sección 9): ninguno de estos arcos se usa en la base, "
   "−15 % ni −30 %, cuyas asignaciones no superan 90 km.")

# %% [markdown]
# **Verificación adicional (5-oct-2026): pares OSRM entre 140 y 160 km.** Por debajo de 160 km la matriz usa OSRM. Como la razón
# Google/OSRM llegó a 1,14 en la verificación del 1-oct (Villamaría–El Cerrito, 222,8 → 254 km), un par OSRM de 158–160 km podría,
# en principio, superar 180 km en Google Maps. Hay **49 pares OSRM en [140, 160) km**; uno (Buenaventura–Palmira) ya estaba
# verificado como control, así que se consultaron los **48 restantes** con la misma metodología (modo carro, primera ruta sugerida,
# URL por par). Resultados en `data/raw/google_maps/verificacion_google_maps_140_160km.csv`.
#
# **Regla:** estos 48 valores **no reemplazan** a OSRM en la matriz final, porque 12 de ellos reflejan cierres viales temporales del
# 5-oct-2026 (la ruta sugerida "evita corte de carretera") y la consulta es de otra fecha que la de los 164 pares del 1-oct. Se usan
# para (1) comprobar que ningún par cambia de factibilidad y (2) construir una matriz alternativa
# (`distancias_km_alt_google_140_160.csv`) que el notebook 02 usa como sensibilidad.

# %%
gm2 = pd.read_csv(RAW / "google_maps" / "verificacion_google_maps_140_160km.csv")
assert not gm2.duplicated(["id_demanda", "id_sitio"]).any() and len(gm2) == 48
assert not set(zip(gm2.id_demanda, gm2.id_sitio)) & set(zip(gm.id_demanda, gm.id_sitio)), "Par repetido con la verificación del 1-oct"
for r in gm2.itertuples():
    assert abs(dist_osrm.loc[r.id_demanda, r.id_sitio] - r.km_OSRM) < 1e-3
franja = [(i, j) for i in dist_osrm.index for j in dist_osrm.columns if 140 <= dist_osrm.loc[i, j] < 160]
no_verif = [p for p in franja if p not in set(zip(gm.id_demanda, gm.id_sitio)) | set(zip(gm2.id_demanda, gm2.id_sitio))]
assert len(franja) == 49 and not no_verif, "Quedan pares de la franja 140–160 km sin verificar"
r2 = gm2.km_Google_Maps / gm2.km_OSRM
cambia2 = int(((gm2.km_OSRM <= D_MAX) != (gm2.km_Google_Maps <= D_MAX)).sum())
md(f"**Resultado:** Google/OSRM mediana {es(r2.median(), 3)}, rango {es(r2.min(), 3)}–{es(r2.max(), 3)}; máximo Google "
   f"{es(gm2.km_Google_Maps.max())} km ({gm2.loc[gm2.km_Google_Maps.idxmax(), 'municipio_origen']}–{gm2.loc[gm2.km_Google_Maps.idxmax(), 'sitio_destino']}). "
   f"**{cambia2} pares cambian de factibilidad a 180 km**: el conjunto de arcos factibles no cambia. "
   f"{gm2.observacion_ruta.str.contains('cierre').sum()} consultas muestran una ruta que evita cierres temporales.")
dist_alt = dist_km.copy()
for r in gm2.itertuples():
    dist_alt.loc[r.id_demanda, r.id_sitio] = r.km_Google_Maps
assert ((dist_alt <= D_MAX) == (dist_km <= D_MAX)).all().all()

# %% [markdown]
# ## 8. Salidas
# **CSV procesados** (`data/processed/`), **base de datos Excel actualizada** (conserva las fórmulas originales del grupo y agrega
# coordenadas, distancias, alcance, fuentes y registro de cambios) y **figura de nodos**.

# %%
D_ids = list(D_nodes.id)
demanda_out = dem.assign(id="D" + dem.divipola).merge(nodos[["id", "lat", "lon"]], on="id")
cols_d = ["id", "divipola", "municipio", "departamento", "lat", "lon"] + [c for c in dem.columns if c not in ("divipola", "municipio", "departamento")]
demanda_out[cols_d].to_csv(PROC / "demanda.csv", index=False, encoding="utf-8")
cand.merge(nodos[["id", "lat", "lon"]], on="id").to_csv(PROC / "candidatos.csv", index=False, encoding="utf-8")
nodos.to_csv(PROC / "nodos_coordenadas.csv", index=False, encoding="utf-8")
dist_km.round(3).to_csv(PROC / "distancias_km.csv", encoding="utf-8")
dist_osrm.round(3).to_csv(PROC / "distancias_km_osrm.csv", encoding="utf-8")
dist_alt.round(3).to_csv(PROC / "distancias_km_alt_google_140_160.csv", encoding="utf-8")   # solo sensibilidad (sección 7)
alb_out = alb[["municipio", "departamento", "albergados", "fecha_dato", "dias_desde_sismo", "tipo_dato", "personas_en_NH",
               "tasa_observada", "veces_f_base", "fuente_id", "fuente", "url", "verificacion", "nota"]]
alb_out.round(5).to_csv(PROC / "tasa_albergue_observada.csv", index=False, encoding="utf-8")
largo = dist_km.rename_axis(index="id_demanda", columns="id_sitio").stack().rename("km").reset_index()
largo["km_OSRM"] = [dist_osrm.loc[i, j] for i, j in zip(largo.id_demanda, largo.id_sitio)]
largo["fuente_km"] = [fuente_km.loc[i, j] for i, j in zip(largo.id_demanda, largo.id_sitio)]
largo["municipio_origen"] = largo.id_demanda.map(nom); largo["sitio_destino"] = largo.id_sitio.map(nom)
largo["km_linea_recta"] = [lin.loc[i, j] for i, j in zip(largo.id_demanda, largo.id_sitio)]
largo["razon_carretera_recta"] = largo.km / largo.km_linea_recta
largo["factible_180km"] = (largo.km <= D_MAX).astype(int)
largo["costo_transporte_persona_COP"] = largo.km * TARIFA
largo["fuente"] = np.where(largo.fuente_km == "Google Maps", "Google Maps (modo carro), consulta 1-oct-2026",
                          "OSRM table v1 driving (OpenStreetMap), consulta 30-sep-2026")
largo.round(3).to_csv(PROC / "distancias_largo.csv", index=False, encoding="utf-8")
pd.concat([por_demanda.assign(nivel="demanda"), por_sitio.assign(nivel="sitio")]).to_csv(PROC / "resumen_alcance.csv", encoding="utf-8")
par_out = par_tab.copy(); par_out["Valor"] = par_out.Valor.astype(str)
par_out.to_csv(PROC / "parametros.csv", index=False, encoding="utf-8")
control_pob.to_csv(PROC / "control_poblacion_DANE.csv", index=False, encoding="utf-8")
json.dump({"presupuesto_base_COP": B_BASE, "reduccion_A": RED_A, "reduccion_B": RED_B, "distancia_max_km": D_MAX,
           "tarifa_COP_km_persona": TARIFA, "costo_variable_COP_persona": V_KIT,
           "costo_kit_alimentacion_COP": KIT_ALIMENTACION, "costo_kit_aseo_COP": KIT_ASEO,
           "ajuste_redondeo_kits_COP": AJUSTE_REDONDEO_KITS,
           "nota_costo_variable": "v = 60.000 COP es el redondeo declarado de 29.730 (alimentación) + 29.730 (aseo) = 59.460 COP",
           "fraccion_base": F_FRAC["d_f10"],
           "tasa_albergue_observada_min": float(TASA_MIN), "tasa_albergue_observada_max": float(TASA_MAX),
           "criterio_ciudad_cercana": f"al menos un municipio afectado a <= {D_MAX:.0f} km por carretera (matriz final)",
           "fuente_distancias": f"Por carretera: Google Maps en {int((fuente_km == 'Google Maps').values.sum())} pares verificados (1-oct-2026) y OSRM/OpenStreetMap en el resto (30-sep-2026)",
           "pares_google_maps": int((fuente_km == "Google Maps").values.sum()), "pares_osrm": int((fuente_km == "OSRM").values.sum()),
           "fuente_coordenadas": "DIVIPOLA DANE (datos.gov.co gdxc-w37w), 30-sep-2026"},
          open(PROC / "metadatos_instancia.json", "w"), indent=2, ensure_ascii=False)
print("CSV procesados:", sorted(p.name for p in PROC.glob("*.csv")))

# %%
# Base de datos Excel actualizada: se parte del libro original (fórmulas intactas) y se agregan hojas
wb = load_workbook(RAW / "base_datos_caso2_original.xlsx")
HDR, FILL = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E78")
AMAR = PatternFill("solid", fgColor="FFF2CC")

def hoja(nombre, df, titulo, nota=None, anchos=None):
    if nombre in wb.sheetnames:
        del wb[nombre]
    w = wb.create_sheet(nombre)
    w["A1"] = titulo; w["A1"].font = Font(bold=True, size=12)
    if nota:
        w["A2"] = nota
    for k, c in enumerate(df.columns, 1):
        cell = w.cell(row=4, column=k, value=str(c)); cell.font = HDR; cell.fill = FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for r, fila in enumerate(df.itertuples(index=False), 5):
        for k, v in enumerate(fila, 1):
            if isinstance(v, (np.floating, float)) and np.isnan(v):
                v = None
            w.cell(row=r, column=k, value=v.item() if hasattr(v, "item") else v)
    w.freeze_panes = "B5"
    for k, a_ in enumerate(anchos or [14] * len(df.columns), 1):
        w.column_dimensions[w.cell(row=4, column=k).column_letter].width = a_
    return w

wb["Candidatos"]["A2"] = ("Candidatos pequeños (10 de 17): decisión adoptada por el grupo (excluirlos es solo una sensibilidad del notebook 02). "
                          "Coordenadas DIVIPOLA (30-sep-2026) en la hoja Coordenadas; matriz final de distancias (OSRM + Google Maps en 164 pares) en Distancias_km.")
hoja("Coordenadas", nodos[["id", "divipola", "municipio", "departamento", "tipo", "rol", "lat", "lon", "fuente_coord"]],
     "Coordenadas de cabeceras municipales (DIVIPOLA, DANE)", anchos=[9, 10, 22, 16, 11, 11, 11, 11, 60])
dk = dist_km.round(1).rename_axis("id_demanda").reset_index()
dk.insert(1, "municipio", dk.id_demanda.map(nom))
dk.columns = ["id_demanda", "municipio"] + [f"{j} {nom[j]}" for j in dist_km.columns]
w = hoja("Distancias_km", dk, "Distancias por carretera d_ij (km) — MATRIZ FINAL usada por el modelo",
     "Celdas amarillas: valor de Google Maps (1-oct-2026). Resto: OSRM/OpenStreetMap (30-sep-2026). Costo de transporte = 500 COP × km × persona.",
     anchos=[11, 20] + [11] * len(dist_km.columns))
for r_, i in enumerate(dist_km.index, 5):
    for c_, j in enumerate(dist_km.columns, 3):
        if fuente_km.loc[i, j] == "Google Maps":
            w.cell(row=r_, column=c_).fill = AMAR
do = dist_osrm.round(1).rename_axis("id_demanda").reset_index()
do.insert(1, "municipio", do.id_demanda.map(nom))
do.columns = ["id_demanda", "municipio"] + [f"{j} {nom[j]}" for j in dist_osrm.columns]
hoja("Distancias_OSRM", do, "Distancias OSRM originales (km) — solo referencia; el modelo usa la hoja Distancias_km",
     anchos=[11, 20] + [11] * len(dist_osrm.columns))
hoja("Alcance_180km", por_sitio.reset_index().rename(columns={"index": "id"}), "Alcance de cada sitio a ≤ 180 km",
     anchos=[8, 22, 11, 14, 18, 12])
ctrl_x = control[["id_demanda", "id_sitio", "municipio_origen", "sitio_destino", "grupo", "km_OSRM", "km_Google_Maps", "km_linea_recta",
                  "diferencia_km", "diferencia_pct", "factible_OSRM", "factible_Google", "cambia_factibilidad", "revisar_>10%",
                  "fecha_consulta", "url_consulta"]]
w = hoja("Control_distancias", ctrl_x, "Verificación OSRM vs. Google Maps — COMPLETADA (164 pares, 1-oct-2026)",
         "Google Maps, modo carro, primera ruta sugerida. En la matriz final manda Google Maps en estos pares. Filas rojas: cambian la factibilidad a 180 km.",
         anchos=[10, 8, 20, 20, 18, 10, 12, 11, 11, 11, 9, 9, 11, 9, 11, 60])
ROJO_F = PatternFill("solid", fgColor="F8CBAD")
for k_, r in enumerate(ctrl_x.itertuples(), 5):
    if r.cambia_factibilidad == 1:
        for c_ in range(1, len(ctrl_x.columns) + 1):
            w.cell(row=k_, column=c_).fill = ROJO_F
U_DANE = "https://www.dane.gov.co/index.php/estadisticas-por-tema/demografia-y-poblacion/proyecciones-de-poblacion"
fuentes = pd.DataFrame([
    ("F01", "UNGRD — Registro Único de Damnificados (RUD) del terremoto del 10-ago-2026, consolidado por municipio, corte 17-sep-2026, "
            "consultado a través del agregador secundario datosdelterremoto.org", "https://datosdelterremoto.org",
     "Dato oficial vía agregador secundario", "NH, familias y personas registradas (demanda d_i y filtros S1, S2, S7)",
     "No verificada en esta revisión: el acceso al sitio no se pudo confirmar el 5-oct-2026; los valores vienen del libro original del grupo"),
    ("F02", "DANE (2025). Proyecciones de población municipal por área 2018-2042 (PPED), actualizado el 30-jul-2025 "
            "(archivo PPED-AreaMun-2018-2042_VP.xlsx en data/raw/dane)", U_DANE, "Dato oficial",
     "Población 2026 de los 47 municipios (47/47 coinciden con el archivo) y categoría de los candidatos",
     "Verificada (5-oct-2026): la página del DANE publica la serie municipal por área 2018-2042 (publicación 8-ago-2025); el archivo dice 'Actualizado el 30 de julio de 2025'"),
    ("F03", "DANE. Proyecciones municipales de población 2020-2035, actualización post-COVID-19 (DCD) "
            "(archivo DCD-area-proypoblacion-Mun-2020-2035-ActPostCOVID-19.xlsx en data/raw/dane)", U_DANE, "Dato oficial",
     "Solo control: sensibilidad de la categoría de los candidatos a la fuente de población",
     "Parcial: el archivo está en data/raw/dane, pero en la página consultada el 5-oct-2026 no se identificó un enlace con ese nombre"),
    ("F04", "DANE. DIVIPOLA — Códigos de municipios geolocalizados (corte 30-dic-2024), portal datos.gov.co, recurso gdxc-w37w (consulta 30-sep-2026)",
     "https://www.datos.gov.co/Mapas-Nacionales/DIVIPOLA-C-digos-municipios-geolocalizados/gdxc-w37w", "Dato oficial",
     "Coordenadas de la cabecera de los 47 municipios",
     "Verificada (5-oct-2026): metadatos del recurso (atribución DANE, corte 30-dic-2024); respuesta original en data/raw"),
    ("F05", "OpenStreetMap + Project OSRM, servicio table v1, perfil driving (consulta 30-sep-2026)",
     "https://router.project-osrm.org/table/v1/driving/", "Servicio de ruteo (datos abiertos)",
     "Distancias por carretera de los pares no verificados en Google Maps",
     "Verificada: 58 respuestas originales guardadas con su URL en data/raw/osrm; 5 llamadas re-consultadas coincidieron"),
    ("F06", "El Tiempo, blog 'Venga le cuento' (14-ago-2026). Kits de ayuda a damnificados por el terremoto desde $14.650",
     "https://blogs.eltiempo.com/venga-le-cuento/2026/08/14/kits-de-ayuda-a-damnificados-por-el-terremoto-desde-14-650/",
     "Prensa", "Costo de un kit de alimentación (29.730) y uno de aseo (29.730); v = 60.000 COP (redondeo)",
     "No verificada: el sitio responde 403 al acceso automatizado (5-oct-2026); título y fecha identificados en el buscador; "
     "los precios vienen del libro del grupo. Corroboración parcial en F17"),
    ("F07", "El Cronista (02-sep-2026). La Ungrd entregará hasta $3,1 millones a familias afectadas por el terremoto",
     "https://www.cronista.com/colombia/actualidad-co/la-ungrd-entregara-hasta-31-millones-a-familias-afectadas-por-el-terremoto-quienes-pueden-cobrarlo/",
     "Prensa", "Referencia del costo fijo por plaza c_f (supuesto provisional) y de T = 3 meses",
     "Verificada (5-oct-2026): 3 meses; $787.907,25 a $1.050.543 por hogar-mes según categoría del municipio"),
    ("F08", "Universidad de La Sabana. Caso 2. Terremoto en Colombia (enunciado, Prof. Gonzalo Mejía)",
     "docs/Caso_2_Terremoto_en_Colombia_enunciado.pdf (archivo local)", "Enunciado",
     "Capacidades 3.000/1.000/500, tarifa 500 COP/km·persona, 180 km, regla del presupuesto", "Verificada (archivo del curso)"),
    ("F09", "Google Maps, indicaciones en carro, primera ruta sugerida (consultas 1-oct-2026 y 5-oct-2026)",
     "https://www.google.com/maps (URL de cada consulta en data/raw/google_maps y en las hojas Control_distancias y Control_140_160km)",
     "Servicio de ruteo", "Distancias de 164 pares de la matriz final (1-oct) y control de 48 pares de 140–160 km (5-oct)",
     "Verificada: cada par tiene su URL; Google muestra km enteros desde 100 km"),
    ("F10", "Association of Bay Area Governments — ABAG (2017). Bay Area Earthquake Shelter Needs (white paper; estimaciones con la metodología Hazus de FEMA)",
     "https://files.mtc.ca.gov/library/pub/ABAG/30232.pdf", "Informe técnico",
     "Sustenta f = 10 % como escenario de planeación y la sensibilidad 8,3 / 11,6 / 13 %",
     "Verificada (5-oct-2026): 11,6 % de los desplazados busca refugio en el escenario Hayward; 8,3 % (Napa) a 13,2 % (Solano) por condado; demografía de EE. UU."),
    ("F11", "OPS/OMS (10-ago-2026). Informe de situación 2: Colombia — Terremoto agosto 2026",
     "https://www.paho.org/es/documentos/informe-situacion-2-colombia-terremoto-agosto-2026-10-agosto-2026", "Informe de situación",
     "Contexto del evento", "Verificada (5-oct-2026): magnitud 7,4; profundidad 103 km; sentido en 16 departamentos y afectaciones en 12"),
    ("F12", "DNP (2014). CONPES 3819: Política nacional para consolidar el Sistema de Ciudades en Colombia",
     "https://colaboracion.dnp.gov.co/CDT/Conpes/Econ%C3%B3micos/3819.pdf", "Documento de política",
     "Umbral de 100.000 habitantes para ciudad intermedia (P12)",
     "NO verificada: el sitio del DNP bloquea el acceso automatizado (5-oct-2026). Se usa F13 como fuente secundaria"),
    ("F13", "Wikipedia. Aglomeraciones urbanas de Colombia", "https://es.wikipedia.org/wiki/Aglomeraciones_urbanas_de_Colombia",
     "Fuente secundaria (enciclopedia)", "Respaldo secundario del umbral de 100.000 habitantes del Sistema de Ciudades",
     "Verificada (5-oct-2026): el Sistema de Ciudades del DNP incluye municipios y aglomeraciones de más de 100.000 habitantes"),
    ("F14", "El Nuevo Siglo (18-ago-2026). 2.013 personas permanecen en albergues tras terremoto en Colombia",
     "https://www.elnuevosiglo.com.co/nacion/2013-personas-permanecen-en-albergues-tras-terremoto-en-colombia", "Prensa",
     "Albergados observados: Pereira 1.300, Manizales 195, Cali 192, Armenia 135",
     "Verificada (5-oct-2026)"),
    ("F15", "El Blog del Ministro (29-sep-2026). Dosquebradas unifica alojamientos temporales en el Centro Vida Argemiro Cárdenas",
     "https://www.elblogdelministro.com/2026/09/dosquebradas-unifica-alojamientos.html", "Prensa (blog)",
     "Albergados observados en Dosquebradas: promedio de 400 personas en 48 días; 45 personas el 27-sep",
     "Verificada (5-oct-2026): se abrió la página (antes solo se tenía el extracto del buscador)"),
    ("F16", "UNGRD (02-sep-2026). UNGRD activa apoyo económico temporal para familias damnificadas por el terremoto",
     "https://portal.gestiondelriesgo.gov.co/Paginas/Noticias/2026/UNGRD-activa-apoyo-economico-temporal-para-familias-damnificadas-por-el-terremoto.aspx",
     "Comunicado oficial", "Corrobora F07: montos por hogar-mes por categoría de municipio, 3 meses",
     "Verificada (5-oct-2026)"),
    ("F17", "Valora Analitik (13-ago-2026). Así puede donar kit de alimentos o aseo con D1 para damnificados del terremoto",
     "https://www.valoraanalitik.com/asi-puede-donar-kit-de-alimentos-o-aseo-con-d1-para-damnificados-del-terremoto/", "Prensa",
     "Corroboración del orden de magnitud del costo de los kits (no reemplaza a F06)",
     "Verificada (5-oct-2026): kits de alimentos o aseo de $15.000 y $30.000"),
    ("F18", "UNGRD (2024). Protocolo de alojamientos temporales (Estrategia Nacional de Respuesta)",
     "https://portal.gestiondelriesgo.gov.co/Documents/ENRE/5-Alojamientos-temporales.pdf", "Documento técnico oficial",
     "Búsqueda de una referencia pública para c_f: el protocolo no publica costos; sí el estándar de 3,5 m² por persona",
     "Verificada (5-oct-2026): sin cifras de costo"),
], columns=["ID", "Fuente", "URL", "Tipo", "Uso en el proyecto", "Verificación"])
assert fuentes.ID.is_monotonic_increasing and fuentes.ID.is_unique and fuentes.URL.str.len().gt(0).all()
fuentes.to_csv(PROC / "fuentes.csv", index=False, encoding="utf-8")   # también la lee el módulo de calidad de datos
hoja("Fuentes", fuentes, "Fuentes de datos (IDs ordenados; estado de verificación al 5-oct-2026)",
     "Una fuente que no se pudo abrir se marca 'No verificada'. Referencias completas en docs/bibliografia.md.", anchos=[6, 70, 60, 22, 50, 60])
cambios = pd.DataFrame([
    ("Coordenadas", "Pendientes (notebook 01 sin ejecutar)", "DIVIPOLA oficial para los 47 nodos", "Requisito del caso; un punto (cabecera municipal) por municipio"),
    ("Distancias", "No existían (sensibilidad preliminar usaba distancias no documentadas)", "Matriz 29×18 por carretera (OSRM + Google Maps en 164 pares)", "Requisito del caso; cifras preliminares quedan reemplazadas"),
    ("Distancias (verificación)", "Pares de control 'PENDIENTE' de cotejar con Google Maps", "164 pares verificados en Google Maps; manda Google Maps en ellos",
     "Mediana Google/OSRM ≈ 1,005; 7 pares cambian de factibilidad a 180 km"),
    ("La Dorada (C17)", "Marcada sin ningún municipio a ≤ 180 km (OSRM 212,9 km)", "Alcanza a Manizales (166 km) y Villamaría (165 km) según Google Maps",
     "Error de ruteo OSRM (rodeo por el sur); 17 candidatos útiles"),
    ("Sensibilidad c_f", "Tabla 'techo' preliminar", "Recalculada con el MILP y distancias reales (notebook 02)", "La tabla preliminar no es reproducible"),
    ("Costo variable", "60.000 COP", "Kit de alimentación 29.730 + kit de aseo 29.730 = 59.460; v = 60.000 se declara como redondeo (+540)", "Transparencia (5-oct-2026: parámetros separados)"),
    ("Fuentes", "Hoja con 9 fuentes sin URL ni estado", "18 fuentes con URL, tipo, uso y verificación; IDs ordenados", "Revisión 5-oct-2026"),
    ("Albergados observados", "Rango observado citado sin tabla ni fuente", "Tabla con fuente: 0,4 % (Cali) a 4,4 % (Dosquebradas) de las personas en NH", "Revisión 5-oct-2026"),
    ("Distancias 140–160 km", "48 pares OSRM sin verificar", "Verificados en Google Maps (5-oct-2026); ninguno cambia de factibilidad", "Revisión 5-oct-2026; se usan como sensibilidad"),
    ("Criterio 'ciudad cercana'", "Implícito", "≥ 1 municipio afectado a ≤ 180 km por carretera, reportado por candidato", "Revisión 5-oct-2026"),
], columns=["Elemento", "Antes", "Ahora", "Motivo"])
hoja("Registro_cambios", cambios, "Registro de correcciones y complementos (30-sep, 1-oct y 5-oct-2026)", anchos=[22, 45, 60, 50])

# Albergados observados y verificación 140–160 km
hoja("Albergados_observados", alb_out.assign(tasa_observada_pct=lambda d_: (100 * d_.tasa_observada).round(2)).drop(columns="tasa_observada"),
     "Tasa de albergue observada = albergados ÷ personas en viviendas no habitables (NH × hogar RUD)",
     "Fechas no comparables (18-ago puntual vs. promedio de 48 días). f = 10 % es un escenario de planeación, no una tasa observada.",
     anchos=[16, 16, 11, 22, 10, 30, 13, 9, 9, 22, 60, 26, 50, 13])
hoja("Control_140_160km", gm2.drop(columns=["metodo"]), "Verificación adicional en Google Maps — pares OSRM de 140 a 160 km (48 pares, 5-oct-2026)",
     "No reemplazan a OSRM en la matriz final (12 consultas muestran cierres viales temporales); se usan como control y sensibilidad.",
     anchos=[10, 8, 20, 20, 10, 10, 26, 50, 60, 11, 9, 9, 9])

# Parámetros: se agregan los dos componentes del costo variable (la fila P03 conserva el valor 60.000 que usan las fórmulas)
wp = wb["Parametros"]
fila = wp.max_row + 1
for k_, (pid, nombre, valor, unidad, tipo, fuente_) in enumerate([
        ("P25", "Costo kit de alimentación (por persona)", KIT_ALIMENTACION, "COP/persona", "Real (prensa)", "F06: El Tiempo 14-ago-2026 (precio minorista D1)"),
        ("P26", "Costo kit de aseo (por persona)", KIT_ASEO, "COP/persona", "Real (prensa)", "F06: El Tiempo 14-ago-2026 (precio minorista D1)"),
        ("P27", "Ajuste por redondeo de v (P03 − P25 − P26)", None, "COP/persona", "Calculado", "P03 = 60.000 es el redondeo declarado de 59.460 (+0,9 %)")]):
    for c_, v_ in enumerate((pid, nombre, valor, unidad, tipo, fuente_), 1):
        wp.cell(row=fila + k_, column=c_, value=v_)
wp.cell(row=fila + 2, column=3, value=f"=C7-C{fila}-C{fila + 1}")
assert wp.cell(row=7, column=1).value == "P03"
wp.cell(row=7, column=6, value=str(wp.cell(row=7, column=6).value) + " La suma es 59.460: v = 60.000 es un redondeo declarado (+540, ver P27).")

# Calidad de datos (tools/calidad_datos.py): PASS / WARNING / FAIL. Objetivo: 0 FAIL; los WARNING son debilidades declaradas.
calidad = evaluar_calidad(ROOT)
calidad.to_csv(TAB / "calidad_datos.csv", index=False, encoding="utf-8")
w = hoja("Calidad_datos", calidad, "Calidad de datos — PASS / WARNING / FAIL (generada por tools/calidad_datos.py)",
         "FAIL = rompe un requisito del enunciado o una regla de construcción (objetivo: 0). WARNING = debilidad conocida y declarada.",
         anchos=[6, 13, 58, 12, 11, 110])
COLOR_Q = {"PASS": "C6EFCE", "WARNING": "FFEB9C", "FAIL": "FFC7CE"}
for k_, r_ in enumerate(calidad.itertuples(), 5):
    w.cell(row=k_, column=5).fill = PatternFill("solid", fgColor=COLOR_Q[r_.resultado])
RES_Q = resumen_calidad(calidad)
assert RES_Q["FAIL"] == 0, calidad[calidad.resultado == "FAIL"]

# Diccionario de datos (archivo, columna, unidad, tipo) — también se escribe en docs/diccionario_datos.md
R, CA, S, ID_ = "Real", "Calculado", "Supuesto", "Identificador"
DICC = {
 "data/processed/demanda.csv": [
  ("id", "Identificador del punto de demanda (D + DIVIPOLA)", "—", ID_), ("divipola", "Código DIVIPOLA del municipio", "—", ID_),
  ("municipio", "Nombre del municipio afectado", "—", ID_), ("departamento", "Departamento", "—", ID_),
  ("lat", "Latitud de la cabecera municipal (F04)", "grados", R), ("lon", "Longitud de la cabecera municipal (F04)", "grados", R),
  ("pob_DANE_2026", "Población proyectada 2026 (F02)", "habitantes", R), ("NH_RUD", "Viviendas no habitables registradas en el RUD (F01)", "viviendas", R),
  ("familias_RUD", "Familias registradas en el RUD (F01)", "familias", R), ("personas_RUD", "Personas registradas en el RUD (F01)", "personas", R),
  ("hogar_RUD", "Tamaño de hogar h_i = personas_RUD / familias_RUD", "personas/familia", CA),
  ("personas_en_NH", "Personas en viviendas no habitables = round(NH_RUD · h_i) (una familia por vivienda)", "personas", CA),
  ("d_f02", "Demanda con f = 2 %: round(NH_RUD · h_i · 0,02)", "personas", CA), ("d_f05", "Demanda con f = 5 %", "personas", CA),
  ("d_f10", "Demanda BASE d_i con f = 10 % (supuesto de planeación, F10)", "personas", CA), ("d_f13", "Demanda con f = 13 %", "personas", CA),
  ("d_f10_hogar3", "Demanda con f = 10 % y hogar de 3,0 personas (sensibilidad)", "personas", CA),
  ("pct_registrado", "Personas registradas en el RUD / población DANE × 100", "%", CA),
  ("calidad_dato_RUD", "Calidad del dato: F fuerte, P parcial, R solo RUD (criterio del grupo)", "—", S)],
 "data/processed/candidatos.csv": [
  ("id", "Identificador del sitio (C01–C17 candidatos, R01 reserva)", "—", ID_), ("divipola", "Código DIVIPOLA", "—", ID_),
  ("municipio", "Municipio del sitio", "—", ID_), ("departamento", "Departamento", "—", ID_),
  ("rol", "Candidato o Reserva (La Tebaida no cumple S7)", "—", S), ("pob_DANE_2026", "Población proyectada 2026 (F02)", "habitantes", R),
  ("NH_RUD", "Viviendas no habitables en el municipio del sitio (F01)", "viviendas", R), ("personas_RUD", "Personas registradas en el RUD (F01)", "personas", R),
  ("pct_registrado", "Personas registradas / población × 100", "%", CA),
  ("categoria", "Grande (≥ 200.000 hab.), Intermedia (≥ 100.000) o Pequeña (cortes P11–P12)", "—", CA),
  ("capacidad_K", "Capacidad K_j según categoría (enunciado: 3.000 / 1.000 / 500)", "personas", R),
  ("factor_escala", "Factor φ de economías de escala: 0,85 / 1,00 / 1,15", "—", S),
  ("costo_fijo_por_plaza_T", "c_f · T · φ = costo fijo por plaza en 3 meses", "COP/plaza", CA),
  ("costo_fijo_fj", "Costo fijo F_j = K_j · c_f · T · φ (c_f = 200.000 COP/plaza-mes, supuesto)", "COP", CA),
  ("costo_variable_vj", "Costo variable v = 60.000 (redondeo de 29.730 + 29.730)", "COP/persona", S),
  ("filtros_S2_S7", "Cumple NH < 250 y % registrado < 5 % (criterio del grupo)", "—", CA),
  ("lat", "Latitud de la cabecera (F04)", "grados", R), ("lon", "Longitud de la cabecera (F04)", "grados", R)],
 "data/processed/distancias_largo.csv": [
  ("id_demanda", "Punto de demanda", "—", ID_), ("id_sitio", "Sitio", "—", ID_),
  ("km", "Distancia por carretera de la matriz final (Google Maps en 164 pares, OSRM en el resto)", "km", R),
  ("km_OSRM", "Distancia OSRM original (F05)", "km", R), ("fuente_km", "Fuente del valor km (Google Maps u OSRM)", "—", ID_),
  ("municipio_origen", "Municipio afectado", "—", ID_), ("sitio_destino", "Municipio del sitio", "—", ID_),
  ("km_linea_recta", "Distancia haversine entre cabeceras", "km", CA), ("razon_carretera_recta", "km / km_linea_recta", "—", CA),
  ("factible_180km", "1 si km ≤ 180 (enunciado)", "0/1", CA), ("costo_transporte_persona_COP", "500 · km (enunciado)", "COP/persona", CA),
  ("fuente", "Descripción de la fuente y fecha de consulta", "—", ID_)],
 "data/processed/tasa_albergue_observada.csv": [
  ("municipio", "Municipio", "—", ID_), ("departamento", "Departamento", "—", ID_),
  ("albergados", "Personas en albergues reportadas por la prensa (F14, F15)", "personas", R), ("fecha_dato", "Fecha o periodo del dato", "fecha", R),
  ("dias_desde_sismo", "Días desde el 10-ago-2026", "días", CA), ("tipo_dato", "Corte puntual o promedio del periodo", "—", R),
  ("personas_en_NH", "Personas en viviendas no habitables (demanda.csv)", "personas", CA),
  ("tasa_observada", "albergados / personas_en_NH", "fracción", CA), ("veces_f_base", "f base (0,10) / tasa_observada", "veces", CA),
  ("fuente_id", "ID en la hoja Fuentes", "—", ID_), ("fuente", "Medio y fecha", "—", ID_), ("url", "Enlace", "—", ID_),
  ("verificacion", "Estado de verificación", "—", ID_), ("nota", "Observaciones", "—", ID_)],
 "results/tablas/calidad_datos.csv": [
  ("id", "Código de la prueba (D demanda, C candidatos, P parámetros, G coordenadas y distancias, F fuentes)", "—", ID_),
  ("area", "Grupo de datos revisado", "—", ID_), ("prueba", "Regla que se comprueba", "—", ID_),
  ("nivel_si_no_cumple", "FAIL si rompe un requisito; WARNING si es una debilidad declarada", "—", S),
  ("resultado", "PASS, WARNING o FAIL", "—", CA), ("detalle", "Evidencia de la prueba", "—", CA)],
}
for archivo_, filas_ in DICC.items():
    cols_ = list(pd.read_csv(ROOT / archivo_, nrows=0).columns)
    assert [f_[0] for f_ in filas_] == cols_, f"Diccionario incompleto o desordenado para {archivo_}"
dicc = pd.DataFrame([(a_, *f_) for a_, fs_ in DICC.items() for f_ in fs_], columns=["Archivo", "Columna", "Descripción", "Unidad", "Tipo"])
hoja("Diccionario", dicc, "Diccionario de datos (Real = dato de fuente o del enunciado; Calculado = derivado; Supuesto = decisión del grupo)",
     anchos=[36, 26, 80, 16, 14])
leeme = pd.DataFrame([
    ("Propósito", "Base de datos del Caso 2 (terremoto del 10-ago-2026): demanda, candidatos, costos, coordenadas y distancias por carretera."),
    ("Cómo se genera", "notebooks/01_datos_coordenadas_distancias.ipynb a partir de data/raw/base_datos_caso2_original.xlsx (fórmulas del grupo intactas)."),
    ("Hojas del grupo", "Parametros (fuente única de supuestos; P25–P27 agregados), Candidatos, Presupuesto, Demanda."),
    ("Hojas agregadas", "Coordenadas, Distancias_km (matriz final), Distancias_OSRM, Alcance_180km, Control_distancias, Control_140_160km, "
                        "Albergados_observados, Fuentes, Registro_cambios, Calidad_datos, Diccionario."),
    ("Calidad de datos", "Hoja Calidad_datos: pruebas PASS / WARNING / FAIL sobre los datos de entrada (objetivo 0 FAIL); misma tabla que results/tablas/calidad_datos.csv."),
    ("Tipos de dato", "Real = dato de fuente o del enunciado; Calculado = derivado con fórmula; Supuesto = decisión del grupo (ver hoja Diccionario)."),
    ("Unidades", "Costos en COP; distancias en km por carretera; capacidades y demanda en personas."),
    ("Supuestos clave", "c_f = 200.000 COP/plaza-mes; φ = 0,85/1,00/1,15; f = 10 %; v = 60.000 COP (redondeo de 29.730 + 29.730)."),
    ("Decisiones metodológicas", "17 candidatos con 10 pequeños (sin ellos no se llega a 15); presupuesto sobre el costo total; matriz combinada Google Maps/OSRM. Alternativas solo como sensibilidad."),
    ("Fuentes", "Hoja Fuentes (ID, URL, tipo, uso y verificación) y docs/bibliografia.md."),
], columns=["Campo", "Descripción"])
hoja("LEEME", leeme, "LEEME — base_datos_caso2.xlsx", anchos=[22, 140])
wb.move_sheet("LEEME", offset=-(len(wb.sheetnames) - 1))

# Documentos generados desde las mismas tablas (una sola fuente de verdad)
DOCS = ROOT / "docs"
lin_ = ["# Diccionario de datos", "", "Generado por `notebooks/01_datos_coordenadas_distancias.ipynb` (misma tabla que la hoja `Diccionario` de",
        "`data/processed/base_datos_caso2.xlsx`). **Tipo:** Real = dato de fuente o del enunciado; Calculado = derivado; Supuesto = decisión del",
        "grupo; Identificador = clave o texto descriptivo. Los IDs de fuente (F01…) remiten a `docs/bibliografia.md`.", ""]
for archivo_, g_ in dicc.groupby("Archivo", sort=False):
    lin_ += [f"## `{archivo_}`", "", "| Columna | Descripción | Unidad | Tipo |", "|---|---|---|---|"]
    lin_ += [f"| `{r_.Columna}` | {r_.Descripción} | {r_.Unidad} | {r_.Tipo} |" for r_ in g_.itertuples()] + [""]
(DOCS / "diccionario_datos.md").write_text("\n".join(lin_), encoding="utf-8")
lin_ = ["# Bibliografía y fuentes de datos", "", "Generado por `notebooks/01_datos_coordenadas_distancias.ipynb` (misma tabla que la hoja `Fuentes` de",
        "`data/processed/base_datos_caso2.xlsx`). Estado de verificación al 5-oct-2026: una fuente que no se pudo abrir se marca **No verificada**",
        "y no se le atribuye ninguna cifra que no venga del libro original del grupo.", ""]
for r_ in fuentes.itertuples():
    lin_ += [f"**[{r_.ID}]** {r_[2]}. {r_.URL}  ", f"*Tipo:* {r_.Tipo}. *Uso:* {r_[5]}.  ", f"*Verificación:* {r_[6]}.", ""]
(DOCS / "bibliografia.md").write_text("\n".join(lin_), encoding="utf-8")
wb.save(PROC / "base_datos_caso2.xlsx")
# Opcional: si LibreOffice está instalado, recalcula el libro para guardar también los valores de las fórmulas
import subprocess, tempfile
soffice = shutil.which("soffice") or shutil.which("libreoffice")
if soffice:
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run([soffice, "--headless", "--convert-to", "xlsx", "--outdir", tmp, str(PROC / "base_datos_caso2.xlsx")],
                       capture_output=True, timeout=180)
        recalc = Path(tmp) / "base_datos_caso2.xlsx"
        if recalc.exists():
            wv, wf = load_workbook(recalc, data_only=True), load_workbook(recalc)
            chk = wv["Presupuesto"]["B5"].value
            assert abs(chk - B_BASE) < 1, "El presupuesto del Excel recalculado no coincide"
            assert wv["Parametros"].cell(row=wp.max_row, column=3).value == AJUSTE_REDONDEO_KITS
            # Todas las fórmulas deben tener valor en caché y ninguna puede dar error (#REF!, #DIV/0!, #VALUE!, #NAME?, #N/A)
            n_form = errores_xl = sin_valor = 0
            for ws_ in wf.worksheets:
                for fila_ in ws_.iter_rows():
                    for c_ in fila_:
                        if isinstance(c_.value, str) and c_.value.startswith("="):
                            n_form += 1
                            v_ = wv[ws_.title][c_.coordinate].value
                            sin_valor += v_ is None
                            errores_xl += isinstance(v_, str) and v_.startswith("#")
            assert errores_xl == 0 and sin_valor == 0, f"Excel con {errores_xl} errores y {sin_valor} fórmulas sin valor"
            shutil.copy(recalc, PROC / "base_datos_caso2.xlsx")
            print(f"Libro recalculado con LibreOffice: presupuesto base en Excel = {chk:,.0f} COP (coincide); "
                  f"{n_form} fórmulas, 0 errores, todas con valor en caché")
else:
    print("LibreOffice no disponible: el libro conserva las fórmulas y Excel las calculará al abrirlo.")
md(f"✔ Base de datos actualizada: `data/processed/base_datos_caso2.xlsx` — hojas: {', '.join(wb.sheetnames)}.")
md(f"✔ **Calidad de datos:** {RES_Q['PASS']} PASS · {RES_Q['WARNING']} WARNING · **{RES_Q['FAIL']} FAIL** "
   "(hoja `Calidad_datos` y `results/tablas/calidad_datos.csv`). Los WARNING son debilidades conocidas que se declaran:")
display(calidad[calidad.resultado != "PASS"][["id", "prueba", "resultado", "detalle"]].style.hide(axis="index"))

# %%
fig, ax = plt.subplots(figsize=(7.5, 8.5))
col_dep = {"Valle del Cauca": "#1f77b4", "Risaralda": "#d62728", "Caldas": "#2ca02c", "Quindío": "#9467bd"}
dd = demanda_out
ax.scatter(dd.lon, dd.lat, s=20 + dd.d_f10 / 8, c=dd.departamento.map(col_dep), alpha=0.75, edgecolor="white", label="Municipio afectado (tamaño = demanda)")
cc = nodos[nodos.tipo == "candidato"]
ax.scatter(cc.lon, cc.lat, marker="s", s=60, facecolor="none", edgecolor="#e8590c", linewidth=1.6, label="Sitio candidato")
for r in nodos.itertuples():
    ax.annotate(r.municipio, (r.lon, r.lat), fontsize=6.5, xytext=(3, 2), textcoords="offset points",
                color="#7a2e05" if r.tipo == "candidato" else "#333")
for dpto, c in col_dep.items():
    ax.scatter([], [], c=c, label=dpto)
ax.set_xlabel("Longitud"); ax.set_ylabel("Latitud"); ax.set_aspect("equal")
ax.set_title("Nodos de la instancia: 29 municipios afectados y 18 sitios (17 + reserva)")
ax.legend(fontsize=7, loc="lower left"); ax.grid(alpha=0.25)
plt.tight_layout(); plt.savefig(FIG / "01_nodos_instancia.png", dpi=160); plt.show()

# %% [markdown]
# ## Limitaciones a declarar
# - OSRM y Google Maps usan la red vial **sin cierres ni daños posteriores al sismo** y devuelven la distancia de la **ruta más
#   rápida** sugerida, no necesariamente la más corta. Google Maps redondea a km enteros los trayectos ≥ 100 km.
# - Cada municipio se representa por su **cabecera** (DIVIPOLA); en municipios extensos (Buenaventura, Dagua, El Cairo, Argelia) la
#   población damnificada puede estar lejos de ese punto.
# - La matriz final combina dos fuentes: Google Maps en los 164 pares verificados (todos los de la franja 160–230 km) y OSRM en el resto.
#   Los 48 pares OSRM de 140–160 km se verificaron el 5-oct-2026 (ninguno cambia de factibilidad) y solo se usan como sensibilidad.
#   **Decisión adoptada:** matriz combinada (Google Maps donde se decide la factibilidad); con solo OSRM la red base no cambia.
# - Tres pares que cambian de factibilidad están a ≤ 3 km del corte de 180 km (Armenia–Pradera 178, El Cairo–Buga 181,
#   Caicedonia–Supía 178) y Google muestra km enteros: su factibilidad es frágil (no afecta a la base ni a los recortes).
# - El RUD es un registro abierto (corte 17-sep-2026) y proviene de un agregador secundario; la demanda puede estar subestimada.
