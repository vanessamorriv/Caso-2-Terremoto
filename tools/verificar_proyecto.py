import json, math, re, sys
from pathlib import Path
import pandas as pd
import pulp
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
fallos = []
def chk(cond, msg):
    print(("  ✔ " if cond else "  ✘ ") + msg)
    if not cond:
        fallos.append(msg)
def es(x, d=0):
    return f"{x:,.{d}f}".replace(",", "X").replace(".", ",").replace("X", ".")
def pct(x, d=1):
    return es(100 * x, d) + " %"
def pct_abajo(x, d=1):
    return pct(math.floor(x * 10 ** (d + 2) + 1e-9) / 10 ** (d + 2), d)
TAB = ROOT / "results" / "tablas"

print("1. Archivos y entorno")
esperados = ["README.md", "requirements.txt", ".gitignore", "data/raw/demanda_terremoto.csv", "data/raw/candidatos_costos.csv",
             "data/raw/albergados_observados.csv", "data/raw/base_datos_caso2_original.xlsx", "data/raw/divipola_coordenadas_datosgov.csv",
             "data/raw/osrm/osrm_table_respuestas.csv", "data/raw/google_maps/verificacion_google_maps.csv",
             "data/raw/google_maps/verificacion_google_maps_140_160km.csv",
             "data/processed/base_datos_caso2.xlsx", "data/processed/distancias_km.csv", "data/processed/distancias_km_osrm.csv",
             "data/processed/distancias_km_alt_google_140_160.csv", "data/processed/tasa_albergue_observada.csv",
             "data/processed/demanda.csv", "data/processed/candidatos.csv", "data/processed/resumen_alcance.csv",
             "notebooks/01_datos_coordenadas_distancias.ipynb", "notebooks/02_modelo_milp.ipynb", "results/resultados_modelo.xlsx",
             "results/tablas/resumen_kpis.csv", "results/tablas/registro_solver.csv", "results/tablas/escenarios_presupuesto_alternativo.csv",
             "results/tablas/uso_pares_cambio_factibilidad.csv", "results/tablas/parametros_revision.csv",
             "docs/formulacion.md", "docs/registro_correcciones.md", "docs/uso_ia.md", "docs/bibliografia.md",
             "docs/diccionario_datos.md", "docs/informe_tecnico.md",
             "tools/calidad_datos.py", "results/tablas/calidad_datos.csv", "data/processed/fuentes.csv",
             "docs/explicacion_warnings_calidad.docx"]
for f in esperados:
    chk((ROOT / f).exists(), f)
chk(len(list(TAB.glob("*.csv"))) == 31, f"31 tablas CSV en results/tablas ({len(list(TAB.glob('*.csv')))})")
chk(len(list((ROOT / "results/figuras").glob("*.png"))) == 11, "11 figuras PNG")
whl = list(ROOT.rglob("*.whl"))
chk(not whl, "sin archivos .whl en el repositorio" + (f": {[w.name for w in whl]}" if whl else ""))
chk("*.whl" in (ROOT / ".gitignore").read_text(encoding="utf-8").split(), ".gitignore incluye *.whl")
chk(not (ROOT / "results" / "rapido").exists(), "no hay salidas del modo RAPIDO mezcladas con las oficiales")
req = (ROOT / "requirements.txt").read_text(encoding="utf-8")
chk(re.search(r"^pulp>=2\.8,<4", req, re.M) is not None, "requirements.txt fija pulp>=2.8,<4 (PuLP 4 rompe la API)")
chk(int(pulp.__version__.split(".")[0]) < 4, f"PuLP instalado {pulp.__version__} < 4")

print("2. Notebooks")
textos_nb = {}
for nb in ["01_datos_coordenadas_distancias.ipynb", "02_modelo_milp.ipynb"]:
    d = json.load(open(ROOT / "notebooks" / nb, encoding="utf-8"))
    code = [c for c in d["cells"] if c["cell_type"] == "code"]
    errores = [o for c in code for o in c.get("outputs", []) if o.get("output_type") == "error"]
    sin_ejecutar = [c for c in code if c.get("execution_count") is None and c["source"]]
    absolutas = [c for c in code if re.search(r"(/home/|/root/|/Users/|[A-Z]:\\\\)", "".join(c["source"]))]
    chk(not errores and not sin_ejecutar, f"{nb}: {len(code)} celdas de código ejecutadas sin errores")
    chk(not absolutas, f"{nb}: sin rutas absolutas")
    src = (ROOT / "notebooks" / "src" / nb.replace(".ipynb", ".py")).read_text(encoding="utf-8")
    chk(all("".join(c["source"]).strip() in src for c in code), f"{nb}: celdas idénticas a notebooks/src/{nb.replace('.ipynb', '.py')}")
    textos_nb[nb] = json.dumps(d, ensure_ascii=False)
nb2 = (ROOT / "notebooks/src/02_modelo_milp.py").read_text(encoding="utf-8")
chk(re.search(r"^RAPIDO = False", nb2, re.M) is not None, "notebook 02 entregado con RAPIDO = False (cifras oficiales)")
chk('"pulp>=2.8,<4", "highspy"' in nb2, "notebook 02 instala pulp<4 y highspy si faltan")
chk(all(f"{u} km" in nb2 for u in (50, 100, 150, 180)) and "Justificación" in nb2, "notebook 02 justifica los umbrales 50/100/150/180 km")

print("3. Datos procesados, fuentes y Excel")
dem = pd.read_csv(ROOT / "data/processed/demanda.csv", dtype={"divipola": str}).set_index("id")
cand = pd.read_csv(ROOT / "data/processed/candidatos.csv", dtype={"divipola": str}).set_index("id")
dist = pd.read_csv(ROOT / "data/processed/distancias_km.csv", index_col=0)
dosrm = pd.read_csv(ROOT / "data/processed/distancias_km_osrm.csv", index_col=0)
dalt = pd.read_csv(ROOT / "data/processed/distancias_km_alt_google_140_160.csv", index_col=0)
raw = pd.read_csv(ROOT / "data/raw/demanda_terremoto.csv", dtype={"divipola": str})
chk(dist.shape == (29, 18) and dist.notna().all().all(), "matriz final 29 × 18 completa")
chk((dem.d_f10.values == raw.d_f10.values).all() and dem.d_f10.sum() == 18880, "demanda base = 18.880 (igual al CSV original)")
resp = pd.read_csv(ROOT / "data/raw/osrm/osrm_table_respuestas.csv")
r0 = resp[(resp.origen_idx == 0) & (resp.bloque == 0)].distancias_m.iloc[0].split(",")
chk(abs(float(r0[1]) / 1000 - dosrm.iloc[0, 0]) < 1e-3, "matriz OSRM procesada = respuestas crudas de OSRM")
gm = pd.read_csv(ROOT / "data/raw/google_maps/verificacion_google_maps.csv")
ok_gm = all(abs(dist.loc[r.id_demanda, r.id_sitio] - r.km_Google_Maps) < 1e-6 for r in gm.itertuples())
n_osrm = sum(abs(dist.loc[i, j] - dosrm.loc[i, j]) < 1e-6 for i in dist.index for j in dist.columns)
chk(ok_gm and n_osrm == dist.size - len(gm) + sum(abs(r.km_Google_Maps - r.km_OSRM) < 1e-6 for r in gm.itertuples()),
    f"matriz final = Google Maps en {len(gm)} pares verificados y OSRM en el resto")
J17 = list(cand.index[cand.rol == "Candidato"])
n_arcos = int((dist[J17] <= 180).values.sum())
chk(n_arcos == 308, f"arcos factibles (≤ 180 km) = {n_arcos} (debe ser 308)")
gm2 = pd.read_csv(ROOT / "data/raw/google_maps/verificacion_google_maps_140_160km.csv")
franja = [(i, j) for i in dosrm.index for j in dosrm.columns if 140 <= dosrm.loc[i, j] < 160]
verif = set(zip(gm.id_demanda, gm.id_sitio)) | set(zip(gm2.id_demanda, gm2.id_sitio))
chk(len(franja) == 49 and len(gm2) == 48 and all(p in verif for p in franja), "los 49 pares OSRM de 140–160 km están verificados (48 nuevos + 1 control)")
chk(gm2.url_consulta.str.startswith("https://www.google.com/maps/dir/").all(), "cada par 140–160 km tiene su URL de Google Maps")
chk(((gm2.km_OSRM <= 180) == (gm2.km_Google_Maps <= 180)).all() and gm2.km_Google_Maps.max() < 180, "ningún par de 140–160 km cambia de factibilidad")
chk(((dalt <= 180) == (dist <= 180)).all().all(), "la matriz alternativa (+48 pares Google) tiene los mismos arcos factibles")
alc = pd.read_csv(ROOT / "data/processed/resumen_alcance.csv")
sit = alc[alc.nivel == "sitio"].set_index("municipio")
chk("cumple_criterio_ciudad_cercana" in alc.columns and (sit.loc[sit.rol == "Candidato", "cumple_criterio_ciudad_cercana"] == "Sí").all(),
    "criterio de 'ciudad cercana' reportado y cumplido por los 17 candidatos")
chk(int(sit.loc["La Dorada", "demandas_a_180km"]) == 2 and int(sit.loc["La Dorada", "pares_que_dependen_de_Google"]) == 2,
    "La Dorada alcanza 2 municipios y ambos dependen de Google Maps")
alb = pd.read_csv(ROOT / "data/processed/tasa_albergue_observada.csv").set_index("municipio")
esper = {"Cali": 0.4, "Armenia": 1.0, "Pereira": 2.3, "Manizales": 2.7, "Dosquebradas": 4.4}
chk(all(round(100 * alb.loc[m, "tasa_observada"], 1) == v for m, v in esper.items()), "tasa de albergue observada: 0,4 / 1,0 / 2,3 / 2,7 / 4,4 %")
meta = json.load(open(ROOT / "data/processed/metadatos_instancia.json", encoding="utf-8"))
chk(meta["costo_kit_alimentacion_COP"] == 29730 and meta["costo_kit_aseo_COP"] == 29730 and meta["ajuste_redondeo_kits_COP"] == 540
    and meta["costo_variable_COP_persona"] == 60000, "v = 29.730 (alimentación) + 29.730 (aseo) + 540 (redondeo) = 60.000")

wb = load_workbook(ROOT / "data/processed/base_datos_caso2.xlsx")
wv = load_workbook(ROOT / "data/processed/base_datos_caso2.xlsx", data_only=True)
chk({"LEEME", "Fuentes", "Diccionario", "Albergados_observados", "Control_140_160km", "Registro_cambios"} <= set(wb.sheetnames),
    "Excel con hojas LEEME, Fuentes, Diccionario, Albergados_observados, Control_140_160km y Registro_cambios")
chk(wb.sheetnames[0] == "LEEME", "LEEME es la primera hoja del Excel")
n_f = n_err = n_none = 0
for ws in wb.worksheets:
    for fila in ws.iter_rows():
        for c in fila:
            if isinstance(c.value, str) and c.value.startswith("="):
                n_f += 1
                v = wv[ws.title][c.coordinate].value
                n_none += v is None
                n_err += isinstance(v, str) and v.startswith("#")
chk(n_f > 0 and n_err == 0 and n_none == 0, f"Excel: {n_f} fórmulas, {n_err} errores, {n_none} sin valor en caché")
ws = wb["Fuentes"]
hdr = [c.value for c in ws[4]]
chk(hdr == ["ID", "Fuente", "URL", "Tipo", "Uso en el proyecto", "Verificación"], "hoja Fuentes con columnas ID, Fuente, URL, Tipo, Uso en el proyecto, Verificación")
fu = pd.DataFrame([[c.value for c in fila] for fila in ws.iter_rows(min_row=5) if fila[0].value], columns=hdr)
chk(fu.ID.is_monotonic_increasing and fu.ID.is_unique, f"IDs de Fuentes ordenados ({fu.ID.iloc[0]}–{fu.ID.iloc[-1]})")
chk(fu.URL.notna().all() and (fu.URL.str.len() > 5).all(), "todas las fuentes tienen URL")
texto_f = " ".join(fu.Fuente)
chk(all(k in texto_f for k in ("ABAG", "OPS", "DNP", "DANE", "El Nuevo Siglo", "El Blog del Ministro", "El Tiempo", "El Cronista", "OSRM", "Google Maps")),
    "Fuentes incluye ABAG, OPS, DNP, DANE, El Nuevo Siglo, El Blog del Ministro, El Tiempo, El Cronista, OSRM y Google Maps")
dnp = fu[fu.Fuente.str.contains("CONPES 3819")]
chk(len(dnp) == 1 and dnp.Verificación.iloc[0].startswith("NO verificada"), "CONPES 3819 (DNP) marcado como NO verificada")
chk("files.mtc.ca.gov/library/pub/ABAG/30232.pdf" in " ".join(fu.URL), "ABAG con su URL")
bib = (ROOT / "docs/bibliografia.md").read_text(encoding="utf-8")
chk(all(f"[{i}]" in bib for i in fu.ID) and all(u.split(" ")[0] in bib for u in fu.URL), "docs/bibliografia.md contiene todos los IDs y URL de la hoja Fuentes")
dic = (ROOT / "docs/diccionario_datos.md").read_text(encoding="utf-8")
for archivo in ["data/processed/demanda.csv", "data/processed/candidatos.csv", "data/processed/distancias_largo.csv", "data/processed/tasa_albergue_observada.csv"]:
    cols = pd.read_csv(ROOT / archivo, nrows=0).columns
    chk(f"`{archivo}`" in dic and all(f"`{c}`" in dic for c in cols), f"diccionario cubre todas las columnas de {archivo}")
chk(wb["Diccionario"].max_row > 50, "hoja Diccionario poblada")

print("3b. Calidad de datos (tools/calidad_datos.py)")
sys.path.insert(0, str(ROOT / "tools"))
from calidad_datos import evaluar as evaluar_calidad, resumen as resumen_calidad
cal = evaluar_calidad(ROOT)
cal_csv = pd.read_csv(TAB / "calidad_datos.csv")
rq = resumen_calidad(cal)
chk(cal[["id", "resultado", "detalle"]].equals(cal_csv[["id", "resultado", "detalle"]]), f"calidad recalculada = results/tablas/calidad_datos.csv ({len(cal)} pruebas)")
chk(rq["FAIL"] == 0, f"calidad de datos: {rq['PASS']} PASS · {rq['WARNING']} WARNING · {rq['FAIL']} FAIL (objetivo 0 FAIL)")
chk(set(cal.resultado) <= {"PASS", "WARNING", "FAIL"} and cal.id.is_unique, "resultados válidos e IDs únicos")
wq = wb["Calidad_datos"] if "Calidad_datos" in wb.sheetnames else None
filas_q = [[c.value for c in f] for f in wq.iter_rows(min_row=5) if f[0].value] if wq else []
chk(wq is not None and [f[0] for f in filas_q] == list(cal.id) and [f[4] for f in filas_q] == list(cal.resultado),
    "hoja Calidad_datos del Excel = tabla de calidad")

print("4. Re-soluciones independientes (implementación mínima, CBC)")
def resolver_indep(J, B, solo_fijos=False):
    I = list(dem.index)
    A = [(i, j) for i in I for j in J if dist.loc[i, j] <= 180]
    m = pulp.LpProblem("verif", pulp.LpMaximize)
    y = {j: pulp.LpVariable(f"y_{j}", cat="Binary") for j in J}
    x = {a: pulp.LpVariable(f"x_{a[0]}_{a[1]}", lowBound=0, cat="Integer") for a in A}
    for i in I:
        m += pulp.lpSum(x[a] for a in A if a[0] == i) <= int(dem.loc[i, "d_f10"])
    for j in J:
        m += pulp.lpSum(x[a] for a in A if a[1] == j) <= int(cand.loc[j, "capacidad_K"]) * y[j]
    fijo = pulp.lpSum(cand.loc[j, "costo_fijo_fj"] / 1e6 * y[j] for j in J)
    costo = fijo + pulp.lpSum((60000 + 500 * dist.loc[a]) / 1e6 * x[a] for a in A)
    m += (fijo if solo_fijos else costo) <= B / 1e6
    solver = pulp.PULP_CBC_CMD(msg=False, gapRel=0, timeLimit=300)
    m.setObjective(pulp.lpSum(x.values())); m.solve(solver)
    ok1 = m.status == pulp.LpStatusOptimal and m.sol_status == pulp.LpSolutionOptimal
    Z = round(pulp.value(m.objective))
    m += pulp.lpSum(x.values()) >= Z; m.setObjective(costo); m.sense = pulp.LpMinimize; m.solve(solver)
    ok2 = m.status == pulp.LpStatusOptimal and m.sol_status == pulp.LpSolutionOptimal
    return ok1 and ok2, Z, sorted(cand.loc[j, "municipio"] for j in J if y[j].value() > 0.5), pulp.value(costo) * 1e6
B = 0.5 * cand.loc[J17].costo_fijo_fj.nlargest(10).sum()
chk(abs(B - 3_547_500_000) < 1, f"presupuesto base = {es(B / 1e6, 1)} M; −15 % = {es(0.85 * B / 1e6, 1)} M; −30 % = {es(0.7 * B / 1e6, 1)} M")
chk(es(0.85 * B / 1e6, 1) == "3.015,4" and es(0.7 * B / 1e6, 1) == "2.483,2", "presupuestos −15 % y −30 % = 3.015,4 M y 2.483,2 M")
ok, Z, ab, costo = resolver_indep(J17, B)
kp = pd.read_csv(TAB / "resumen_kpis.csv").set_index("indicador")
chk(ok, "CBC certifica optimalidad en ambas etapas (LpStatus = Optimal y sol_status = LpSolutionOptimal)")
chk(Z == int(kp.loc["Población atendida", "Base"]) == 5736, f"atendidos base = {es(Z)} (coincide con resultados exportados)")
chk(", ".join(ab) == ", ".join(sorted(kp.loc["Sitios abiertos", "Base"].split(", "))) == "Palmira, Tuluá", f"sitios abiertos = {', '.join(ab)}")
chk(abs(costo - float(kp.loc["Costo total", "Base"])) < 1e3, f"costo total = {es(costo / 1e6, 1)} M")
alt = pd.read_csv(TAB / "escenarios_presupuesto_alternativo.csv")
J7 = [j for j in J17 if cand.loc[j, "categoria"] != "Pequeña"]
B7 = 0.5 * cand.loc[J7].costo_fijo_fj.sum()
chk(abs(B7 - 3_030_000_000) < 1, f"B recalculado sin pequeños = 50 % × {es(cand.loc[J7].costo_fijo_fj.sum() / 1e6)} M = {es(B7 / 1e6)} M")
for nivel, fac, z_esp in (("base", 1.0, 4089), ("−15 %", 0.85, 4000), ("−30 %", 0.70, 3000)):
    ok, Z, ab, _ = resolver_indep(J7, B7 * fac)
    fila = alt[(alt.grupo == "Sin pequeños, B recalculado") & (alt.nivel == nivel)].iloc[0]
    chk(ok and Z == z_esp == fila.atendidos and ", ".join(ab) == ", ".join(sorted(fila.sitios.split(", "))),
        f"sin pequeños, B recalculado, {nivel}: {es(Z)} atendidos ({', '.join(ab)}) = notebook")
for nivel, fac, z_esp in (("base", 1.0, 6500), ("−15 %", 0.85, 5000), ("−30 %", 0.70, 4500)):
    ok, Z, ab, _ = resolver_indep(J17, B * fac, solo_fijos=True)
    fila = alt[(alt.grupo == "R4': tope solo sobre costos fijos") & (alt.nivel == nivel)].iloc[0]
    chk(ok and Z == z_esp == fila.atendidos, f"R4' (solo fijos), {nivel}: {es(Z)} atendidos (notebook: {es(fila.atendidos)}; sitios notebook: {fila.sitios}; independiente: {', '.join(ab)})")
fila = alt[(alt.grupo == "Sin pequeños, B conservado") & (alt.nivel == "base")].iloc[0]
chk(fila.atendidos == 5736 and fila.sitios == "Palmira, Tuluá", "sin pequeños con B conservado: la base no cambia (5.736, Palmira y Tuluá)")

print("5. Auditoría del solver")
reg = pd.read_csv(TAB / "registro_solver.csv")
inf = reg[reg.estado == "Infeasible"]
N_OPT = int((reg.estado == "Optimal").sum())
chk(set(reg.estado) <= {"Optimal", "Infeasible"}, f"{len(reg)} llamadas: {N_OPT} óptimas certificadas + {len(inf)} infactible; ninguna cortada por tiempo")
chk(len(inf) == 1 and inf.modelo.iloc[0].startswith("prueba_infactible"), "la única infactible es la prueba intencional (piso de la siguiente décima, 29,1 %)")

print("6. Cifras que no deben cambiar y cifras del README")
readme = (ROOT / "README.md").read_text(encoding="utf-8")
cols = list(kp.columns)
g = lambda ind, c: float(kp.loc[ind, c])
chk([int(g("Población atendida", c)) for c in cols[:3]] == [5736, 4500, 4000], "atendidos base / −15 % / −30 % = 5.736 / 4.500 / 4.000")
chk(abs(g("Costo kits alimentación", "Base") + g("Costo kits aseo", "Base") + g("Costo ajuste por redondeo de kits", "Base")
        - g("Costo atención (kits)", "Base")) < 1, "KPI de kits: alimentación + aseo + redondeo = costo de atención")
par = pd.read_csv(TAB / "extension_parametros.csv").set_index("parametro").valor
ext = pd.read_csv(TAB / "extension_escenarios.csv")
umb = pd.read_csv(TAB / "umbral_distancia.csv")
chk(par["umbral_distancia_km_red_base"] == 90 and not umb.loc[umb.dmax_km == 89, "igual_a_base"].iloc[0], "umbral de distancia 90 km (con 89 km ya cambia la red)")
chk(pct_abajo(par["alpha_max"], 2) == "29,05 %" and abs(par["alpha_infactible_comprobado"] - 0.291) < 1e-9, "α* = 29,05 % (29,1 % infactible)")
chk(pct_abajo(par["beta_max"], 2) == "27,05 %", "β* = 27,05 %")
chk(int(ext.precio_equidad.iloc[0]) == 250, "precio de α* = 250 personas")
frb = pd.read_csv(TAB / "extension_frontera_beta.csv")
fr = pd.read_csv(TAB / "extension_frontera.csv")
b25 = frb[(frb.piso_beta - 0.025).abs() < 1e-4].iloc[0]
chk(int(b25.precio_equidad) == 49 and int(b25.mun_sin_atencion) == 0, "β = 2,5 % cuesta 49 personas y deja 0 municipios en 0 %")
chk(int(fr[(fr.piso_alpha - 0.25).abs() < 1e-9].mun_sin_atencion.iloc[0]) == 22, "con α = 25 % quedan 22 municipios en 0 %")
cf = pd.read_csv(TAB / "sensibilidad_costo_fijo_modelo.csv")
cfb = cf[(cf.regla_presupuesto == "B fijo") & (cf.f == "10%") & (cf.escenario == "base")].set_index("c_f").atendidos_modelo
chk(cfb[100000] == 9767 and cfb[400000] == 3000, "c_f con B fijo: 9.767 (100.000) y 3.000 (400.000)")
uso = pd.read_csv(TAB / "uso_pares_cambio_factibilidad.csv")
chk(len(uso) == 7 and uso.a_3km_del_corte.sum() == 3 and uso.usado_en_base_o_recortes.sum() == 0,
    "7 pares cambian de factibilidad, 3 a ≤ 3 km del corte, ninguno usado en base/−15 %/−30 %")
cifras = [es(g("Población atendida", c)) for c in cols] + [es(g("Población no atendida", c)) for c in cols]
cifras += [es(g("Distancia promedio ponderada (km)", c), 1) for c in cols] + [es(g("Distancia máxima recorrida (km)", c), 1) for c in cols]
cifras += ["$" + es(g("Presupuesto", c) / 1e6, 1) + " M" for c in cols[:3]] + [pct(g("% atendida", c)) for c in cols]
cifras += [es(g("Costo total", c) / 1e6, 1) for c in cols] + [es(g("Costo kits alimentación", c) / 1e6, 1) for c in cols]
cifras += [es(g("Costo ajuste por redondeo de kits", c) / 1e6, 1) for c in cols]
cifras += [es(int(ext.precio_equidad.iloc[0])), pct_abajo(ext.alpha_max.iloc[1], 2), pct_abajo(par["alpha_max"], 2), pct_abajo(par["alpha_max"], 1),
           pct_abajo(par["beta_max"], 2), es(par["umbral_distancia_km_red_base"])]
cifras += [es(int(r.atendidos)) for r in alt.itertuples()] + ["$3.030,0 M", "30,4 %", "29.730", "59.460", "60.000"]
cifras += [f"{len(reg)} llamadas", f"{N_OPT} óptimas + 1 infactible a propósito"]
faltan = [c for c in cifras if c not in readme]
chk(not faltan, f"{len(cifras)} cifras clave presentes en el README" + (f" — faltan: {faltan}" if faltan else ""))
chk(re.search(r"α\s*=\s*29\s*%|piso departamental del 29 %", readme) is None, "el README no redondea α* a 29 % (sería infactible)")
chk("pulp>=2.8,<4" in readme, "el README documenta la restricción de versión de PuLP")
chk("PENDIENTE (grupo)" in readme and "Integrantes" in readme and "Licencia" in readme, "README con integrantes y licencia como PENDIENTE (grupo)")

print("7. Frases desactualizadas")
textos = {p.relative_to(ROOT).as_posix(): p.read_text(encoding="utf-8") for p in
          [ROOT / "README.md", *(ROOT / "docs").glob("*.md"), *(ROOT / "notebooks/src").glob("*.py")]}
textos.update(textos_nb)
for ws_ in wv.worksheets:
    textos[f"Excel:{ws_.title}"] = " ".join(str(c.value) for fila in ws_.iter_rows() for c in fila if c.value is not None)
prohibidas = {"53/53": "validación DIVIPOLA no reproducible (retirada en A12)",
              "cerca del 31 %": "debe ser ≈ 30,4 % (5.736/18.880)",
              "ninguna sin certificar": "hubo 1 infactible a propósito",
              "≈ 630": "usar el número real de llamadas",
              "observado ≈ 2–5 %": "reemplazado por la tabla de albergados observados",
              "Sin pequeños la base no cambia;": "solo vale con B conservado",
              "desaparecen `PULP_CBC_CMD` y `LpStatus`": "afirmación sin el detalle verificado",
              "+ 0,1 pp": "la prueba usa la siguiente décima de punto (29,1 %), no α* + 0,1 pp",
              "una décima de punto mayor": "la prueba usa la siguiente décima de punto",
              "ninguno en un municipio afectado": "los candidatos no son municipios de demanda, pero pueden tener daños menores",
              "Matriz OSRM 29×18": "la matriz final combina OSRM y Google Maps",
              "≈ 15 min": "la ejecución completa tarda ≈ 11 min",
              "pendiente de confirmar con el profesor": "las tres decisiones metodológicas ya son definitivas",
              "pendientes del profesor": "las tres decisiones metodológicas ya son definitivas",
              "sujeta a confirmación del profesor": "las tres decisiones metodológicas ya son definitivas"}
for frase, motivo in prohibidas.items():
    # el registro de correcciones puede citar la frase antigua para documentar el cambio
    norm = lambda s: s.replace("**", "").replace("\\*", "*").lower()   # ignora negritas y escapes de markdown, y mayúsculas
    donde = [k for k, v in textos.items() if norm(frase) in norm(v) and not k.endswith("registro_correcciones.md")]
    chk(not donde, f"sin '{frase}' ({motivo})" + (f" — aparece en {donde}" if donde else ""))

print("\nRESULTADO:", "TODO CORRECTO" if not fallos else f"{len(fallos)} problema(s)")
sys.exit(1 if fallos else 0)
